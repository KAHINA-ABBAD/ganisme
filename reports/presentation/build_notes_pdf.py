"""Genere commentaires_oral.pdf : pour chaque diapositive, sa miniature, l'orateur et le texte a dire.

    .venv/Scripts/python.exe reports/presentation/build_notes_pdf.py <dossier des miniatures s01.png...>
"""

import sys
from pathlib import Path
from xml.sax.saxutils import escape

from pptx import Presentation
from reportlab.lib.colors import HexColor
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import cm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Image, KeepTogether, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

HERE = Path(__file__).resolve().parent
DECK = HERE / "GANisme_sombre.pptx"
OUT = HERE / "commentaires_oral.pdf"
THUMBS = Path(sys.argv[1])

pdfmetrics.registerFont(TTFont("Calibri", "C:/Windows/Fonts/calibri.ttf"))
pdfmetrics.registerFont(TTFont("Calibri-Bold", "C:/Windows/Fonts/calibrib.ttf"))

INK, MUTED, RED, GOLD = HexColor("#1E1B18"), HexColor("#6B645D"), HexColor("#8C2F39"), HexColor("#B8892E")
title_style = ParagraphStyle("t", fontName="Calibri-Bold", fontSize=22, leading=26, textColor=INK)
sub_style = ParagraphStyle("s", fontName="Calibri", fontSize=11, leading=14, textColor=MUTED)
head_style = ParagraphStyle("h", fontName="Calibri-Bold", fontSize=13, leading=16, textColor=INK)
who_style = ParagraphStyle("w", fontName="Calibri-Bold", fontSize=10, leading=12, textColor=RED)
body_style = ParagraphStyle("b", fontName="Calibri", fontSize=12, leading=16.5, textColor=INK, spaceAfter=5)
part_style = ParagraphStyle("p", fontName="Calibri-Bold", fontSize=16, leading=20, textColor=RED, spaceBefore=6)

prs = Presentation(DECK)
story = [
    Paragraph("GANisme · texte de la soutenance", title_style),
    Spacer(1, 4),
    Paragraph("Kahina ABBAD : diapositives 1 à 12 · Pape Ahmed Fall : diapositives 13 à 24 · environ 15 minutes",
              sub_style),
    Spacer(1, 14),
]
current = None
for i, slide in enumerate(prs.slides, 1):
    texts = [sh.text_frame.text for sh in slide.shapes if sh.has_text_frame and sh.text_frame.text.strip()]
    title = texts[0] if i == 1 else texts[1]
    note = slide.notes_slide.notes_text_frame.text
    speaker, _, note = note.partition(" — ")
    if speaker != current:
        current = speaker
        name = "Kahina" if speaker == "KAHINA" else "Pape"
        story.append(Paragraph(f"Partie de {name}", part_style))
        story.append(Spacer(1, 6))
    thumb = Image(str(THUMBS / f"s{i:02d}.png"), width=6.2 * cm, height=6.2 * cm * 9 / 16)
    lines = [Paragraph(escape(line), body_style) for line in note.split("\n") if line.strip()]
    right = [Paragraph(f"{i}. {escape(title)}", head_style), Spacer(1, 2),
             Paragraph(speaker.capitalize(), who_style), Spacer(1, 6)] + lines
    table = Table([[thumb, right]], colWidths=[6.6 * cm, None])
    table.setStyle(TableStyle([
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 0),
        ("RIGHTPADDING", (0, 0), (0, 0), 10),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 12),
        ("LINEBELOW", (0, 0), (-1, -1), 0.5, HexColor("#D9D4CE")),
    ]))
    story += [KeepTogether(table), Spacer(1, 10)]

doc = SimpleDocTemplate(str(OUT), pagesize=A4, leftMargin=1.6 * cm, rightMargin=1.6 * cm,
                        topMargin=1.5 * cm, bottomMargin=1.5 * cm, title="GANisme · texte de la soutenance",
                        author="Kahina ABBAD & Pape Ahmed Fall")


def page_number(canvas, document):
    canvas.setFont("Calibri", 9)
    canvas.setFillColor(MUTED)
    canvas.drawRightString(A4[0] - 1.6 * cm, 0.8 * cm, f"{document.page}")


doc.build(story, onFirstPage=page_number, onLaterPages=page_number)
print(OUT)
