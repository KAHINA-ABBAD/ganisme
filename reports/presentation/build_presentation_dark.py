"""Genere reports/presentation/GANisme_sombre.pptx : version sombre, mots-cles a l'ecran,
explications dans les commentaires (notes du presentateur).

Les images affichees sont choisies a la main (liste SAFE_IMAGES, verifiee visuellement :
pas de nus). Les chiffres viennent des resultats reels du projet.

    .venv/Scripts/python.exe reports/presentation/build_presentation_dark.py
"""

import json
import random
import subprocess
import sys
from datetime import datetime
from pathlib import Path

import numpy as np
import pandas as pd
from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION, XL_MARKER_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.oxml.ns import qn
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "reports" / "presentation" / "GANisme_sombre.pptx"
EDA = json.loads((ROOT / "reports" / "eda" / "summary.json").read_text(encoding="utf-8"))
MANIFEST = pd.read_csv(ROOT / "data" / "processed" / "manifest.csv")
PAINTINGS = pd.read_csv(ROOT / "data" / "processed" / "paintings_metadata.csv")
IMAGES = ROOT / "data" / "processed" / "images_256"
FIG_EDA = ROOT / "reports" / "eda" / "figures"
FIG_DS = ROOT / "reports" / "dataset" / "figures"
FIG_TR = ROOT / "reports" / "training" / "figures"
CHECKS = json.loads((ROOT / "reports" / "dataset" / "checks.json").read_text(encoding="utf-8"))
TRAIN_JSON = ROOT / "reports" / "training" / "metrics.json"
TRAIN = json.loads(TRAIN_JSON.read_text(encoding="utf-8")) if TRAIN_JSON.exists() else {}
TRAIN_DONE = "fid" in TRAIN
IS_V0 = "v0" in TRAIN.get("dataset", "")
MODEL_LABEL = (f"baseline v0 · jeu partiel ({TRAIN['n_train']:,} images)".replace(",", " ") if IS_V0
               else f"modèle final · {TRAIN.get('n_train', 0):,} images".replace(",", " "))
LOG_LINES = (ROOT / "data" / "processed" / "download.log").read_text(encoding="utf-8").splitlines()[-8:]
_tests = subprocess.run([sys.executable, "-m", "pytest", "-v", "-p", "no:cacheprovider", "--no-header"],
                        cwd=ROOT, env={**__import__("os").environ, "PYTHONPATH": str(ROOT / "src")},
                        capture_output=True, text=True)
TEST_LINES = [l.replace("tests/", "").split(" [")[0] for l in _tests.stdout.splitlines() if "PASSED" in l or "FAILED" in l][-12:]
TEST_LINES += [l.strip("= ") for l in _tests.stdout.splitlines() if " passed" in l or " failed" in l][-1:]

# Images verifiees une a une (aucun nu). Par genre, pour composer les visuels.
SAFE_IMAGES = {
    "still-life": [696, 1485, 302, 2050, 1555, 1570, 327, 188, 558, 582, 37, 1090, 547, 2059, 664,
                   1685, 393, 544, 38, 697, 44, 2033, 1433, 2073, 1480, 1027, 1427],
    "landscape": [630, 2088, 1776, 1296, 852, 1837, 1929, 1502, 73, 2053, 1545, 1033, 1508, 1149, 729,
                  1510, 522, 711, 1838, 441, 908, 1514, 1096, 1082, 2081, 2008, 445, 1835],
    "interior": [346, 1995, 348, 1072, 1449, 1736, 1065, 921, 1323, 1996, 869, 1994, 1998, 1548, 349,
                 617, 491, 2000, 1085, 1999, 983, 884, 353, 596, 399, 251, 642],
    "portrait": [686, 2094, 1843, 1253, 886, 1903, 512, 1673, 1623, 78, 1909, 1937, 964, 1576, 2091,
                 1645, 453, 286, 2084, 1906, 358, 945, 1664, 1217, 1211, 963, 701, 374, 1861],
    "genre": [919, 488, 763, 1043, 988, 1265, 1864, 53, 709, 1771, 1984, 1614, 987, 477, 1089, 1423,
              1437, 562, 1550, 1738, 1953, 1266, 635, 1044, 1547],
}

_done = set(MANIFEST.loc[MANIFEST["status"] == "processed", "image_id"])
SAFE_ROWS = [next(i for i in SAFE_IMAGES[g] if i in _done) for g in ("portrait", "landscape", "still-life", "interior", "genre")]

BG = RGBColor(0x14, 0x12, 0x10)
CARD = RGBColor(0x26, 0x22, 0x1E)
CARD2 = RGBColor(0x33, 0x2D, 0x28)
TEXT = RGBColor(0xF2, 0xED, 0xE6)
MUTED = RGBColor(0xA3, 0x9A, 0x90)
GOLD = RGBColor(0xD4, 0xA8, 0x4B)
RED = RGBColor(0xC0, 0x50, 0x5C)
STONE = RGBColor(0x6E, 0x66, 0x5E)
HEAD, BODY = "Cambria", "Calibri"

W, H = Inches(13.333), Inches(7.5)
M = Inches(0.6)
cw = Inches(3.85)  # largeur d'une colonne sur trois
prs = Presentation()
prs.slide_width, prs.slide_height = W, H
BLANK = prs.slide_layouts[6]
GENRE_FR = {
    "religious": "religieux", "portrait": "portrait", "landscape": "paysage",
    "mythological": "mythologique", "genre": "scène de genre", "still-life": "nature morte",
    "other": "autre", "historical": "historique", "interior": "intérieur", "study": "étude",
}


def fmt(n):
    return f"{n:,}".replace(",", " ")


def pct(x):
    return f"{x * 100:.0f} %"


def image_path(image_id):
    path = IMAGES / f"{image_id:05d}.jpg"
    assert path.exists(), f"image verifiee absente : {path}"
    return path


def safe(n, seed, genres=None):
    ids = [i for g in (genres or SAFE_IMAGES) for i in SAFE_IMAGES[g]]
    random.Random(seed).shuffle(ids)
    return [image_path(i) for i in ids[:n]]


# ---------------------------------------------------------------- primitives

def slide_(notes):
    slide = prs.slides.add_slide(BLANK)
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = BG
    slide.notes_slide.notes_text_frame.text = notes
    return slide


def text(slide, x, y, w, h, content, size=18, color=TEXT, bold=False, font=BODY,
         align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, space_after=4):
    x, y, w, h = (int(v) for v in (x, y, w, h))  # PowerPoint refuse les positions decimales
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    for i, para in enumerate(content if isinstance(content, list) else [content]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space_after)
        for run_text, opts in (para if isinstance(para, list) else [(para, {})]):
            run = p.add_run()
            run.text = run_text
            run.font.name = opts.get("font", font)
            run.font.size = Pt(opts.get("size", size))
            run.font.bold = opts.get("bold", bold)
            run.font.color.rgb = opts.get("color", color)
    return box


def shape(slide, kind, x, y, w, h, fill, radius=None):
    s = slide.shapes.add_shape(kind, *(int(v) for v in (x, y, w, h)))
    if radius is not None:
        s.adjustments[0] = radius
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    s.line.fill.background()
    s.shadow.inherit = False
    return s


def card(slide, x, y, w, h, fill=CARD):
    return shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, fill, radius=0.06)


def chip(slide, x, y, w, label, fill=CARD2, color=TEXT, size=16, h=Inches(0.55)):
    """Mot-cle dans une pastille arrondie."""
    s = shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, fill, radius=0.5)
    tf = s.text_frame
    tf.margin_left = tf.margin_right = Inches(0.12)
    tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = label
    run.font.name, run.font.size, run.font.bold = BODY, Pt(size), True
    run.font.color.rgb = color
    return s


def circle(slide, x, y, d, label, fill=RED, size=18):
    s = shape(slide, MSO_SHAPE.OVAL, x, y, d, d, fill)
    tf = s.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = label
    run.font.name, run.font.size, run.font.bold = BODY, Pt(size), True
    run.font.color.rgb = TEXT
    return s


def arrow(slide, x, y, w=Inches(0.35), h=Inches(0.3)):
    return shape(slide, MSO_SHAPE.RIGHT_ARROW, x, y, w, h, STONE)


def line(slide, x1, y1, x2, y2, color=STONE, width=2.5, head=False, dash=False):
    """Trait (eventuellement flechi ou pointille) entre deux points."""
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, *(int(v) for v in (x1, y1, x2, y2)))
    c.line.color.rgb = color
    c.line.width = Pt(width)
    ln = c.line._get_or_add_ln()
    if dash:
        ln.append(ln.makeelement(qn("a:prstDash"), {"val": "dash"}))
    if head:
        ln.append(ln.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"}))
    return c


def dot(slide, cx, cy, d, fill, outline=None):
    s = slide.shapes.add_shape(MSO_SHAPE.OVAL, *(int(v) for v in (cx - d / 2, cy - d / 2, d, d)))
    s.shadow.inherit = False
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if outline is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = outline
        s.line.width = Pt(2.25)
    return s


def label(slide, x, y, w, content, size=16, color=TEXT, bold=True, align=PP_ALIGN.CENTER, h=Inches(0.45)):
    return text(slide, x, y, w, h, content, size=size, color=color, bold=bold, align=align,
                anchor=MSO_ANCHOR.MIDDLE)


def node(slide, x, y, w, h, title, fill=CARD2, size=18, color=TEXT, sub=None):
    """Boite de schema avec un titre (et un sous-titre facultatif)."""
    card(slide, x, y, w, h, fill=fill)
    content = [title] if sub is None else [title, [(sub, {"size": size - 4, "color": GOLD, "bold": False})]]
    text(slide, x, y, w, h, content, size=size, bold=True, color=color, align=PP_ALIGN.CENTER,
         anchor=MSO_ANCHOR.MIDDLE, space_after=0)


def trap(slide, intuition, finding, y=Inches(6.05)):
    """Bandeau DECOUVERTE : l'intuition de depart, puis ce que nos recherches ont montre."""
    card(slide, M, y, Inches(12.13), Inches(0.85), fill=CARD)
    chip(slide, M + Inches(0.2), y + Inches(0.17), Inches(1.75), "DÉCOUVERTE", fill=RED, size=14, h=Inches(0.5))
    text(slide, M + Inches(2.2), y, Inches(4.6), Inches(0.85), [[("Intuition : ", {"color": MUTED}), (intuition, {"color": MUTED})]],
         size=16, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    text(slide, M + Inches(6.9), y, Inches(5.1), Inches(0.85), [[("→ ", {"color": GOLD}), (finding, {})]],
         size=17, bold=True, anchor=MSO_ANCHOR.MIDDLE)


def figure(slide, path, x, y, max_w, max_h, caption=None):
    """Figure reelle du projet, a l'echelle, centree dans la zone, avec sa legende de source."""
    with Image.open(path) as im:
        ratio = im.width / im.height
    cap_h = Inches(0.3) if caption else 0
    w_, h_ = max_w, int(max_w / ratio)
    if h_ > max_h - cap_h:
        h_ = max_h - cap_h
        w_ = int(h_ * ratio)
    px = int(x + (max_w - w_) / 2)
    pic = slide.shapes.add_picture(str(path), px, int(y), int(w_), int(h_))
    if caption:
        text(slide, px, y + h_ + Inches(0.05), w_, Inches(0.25), caption, size=11, color=MUTED)
    return pic


def result_figure(slide, path, x, y, w, h, caption):
    """Resultat d'entrainement s'il existe, sinon un cadre « en cours » (jamais de resultat invente)."""
    if path.exists():
        return figure(slide, path, x, y, w, h, caption)
    card(slide, x, y, w, h, fill=CARD)
    text(slide, x, y, w, h, [[("en cours", {"size": 26, "color": GOLD, "font": HEAD, "bold": True})],
                             [(caption, {"size": 15, "color": MUTED})]],
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


def terminal(slide, x, y, w, h, lines, title):
    """Sortie reelle d'une commande ou d'un journal, dans un cadre facon terminal."""
    card(slide, x, y, w, h, fill=RGBColor(0x0B, 0x0A, 0x09))
    text(slide, x + Inches(0.25), y + Inches(0.15), w - Inches(0.5), Inches(0.3), title, size=12,
         color=GOLD, font="Consolas")
    text(slide, x + Inches(0.25), y + Inches(0.55), w - Inches(0.5), h - Inches(0.7),
         [[(line, {"color": RGBColor(0x9F, 0xD3, 0x9A) if "passed" in line or "PASSED" in line else TEXT})]
          for line in lines], size=9 if max(map(len, lines)) > 80 else 11, font="Consolas", space_after=1)


def header(slide, kicker, title):
    text(slide, M, Inches(0.45), Inches(12), Inches(0.3), kicker.upper(), size=12, bold=True, color=GOLD)
    text(slide, M, Inches(0.75), Inches(12.1), Inches(0.8), title, size=34, bold=True, font=HEAD)


def picture(slide, path, x, y, side):
    return slide.shapes.add_picture(str(path), *(int(v) for v in (x, y, side, side)))


def big_number(slide, x, y, w, value, label, color=GOLD, size=48):
    text(slide, x, y, w, Inches(0.95), value, size=size, bold=True, font=HEAD, color=color,
         align=PP_ALIGN.CENTER)
    text(slide, x, y + Inches(1.0), w, Inches(0.5), label, size=16, color=MUTED, align=PP_ALIGN.CENTER)


def dark_chart(chart, size=13):
    chart.font.name = BODY
    chart.font.size = Pt(size)
    chart.font.color.rgb = MUTED
    chart.has_title = False


def divider(number, title, keywords, notes, seed):
    slide = slide_(notes)
    text(slide, M, Inches(2.3), Inches(4), Inches(1.4), number, size=110, bold=True, font=HEAD, color=GOLD)
    text(slide, M, Inches(3.9), Inches(6.5), Inches(0.9), title, size=44, bold=True, font=HEAD)
    for k, word in enumerate(keywords):
        chip(slide, M + k * Inches(2.05), Inches(5.0), Inches(1.9), word, size=14)
    side, gap = Inches(1.45), Inches(0.1)
    for k, path in enumerate(safe(9, seed)):
        r, c = divmod(k, 3)
        picture(slide, path, Inches(7.95) + c * (side + gap), Inches(1.4) + r * (side + gap), side)
    return slide


# ---------------------------------------------------------------- donnees
status = MANIFEST["status"].value_counts()
n_ok, n_fail = int(status.get("processed", 0)), int(status.get("failed", 0))
n_skip, n_pend = int(status.get("skipped", 0)), int(status.get("pending", 0))
stamp = datetime.now().strftime("%d/%m/%Y à %H:%M")
types = pd.Series(EDA["types"])
n_paint = EDA["n_paintings"]

# ================================================================ 1. Titre
s = slide_(
    "Bonjour. Notre projet, GANisme, consiste à entraîner un réseau antagoniste génératif (GAN) "
    "capable de créer des peintures qui n'existent pas, à partir du catalogue de la Web Gallery of Art.\n\n"
    "Toutes les images de cette présentation sont de vraies peintures de notre jeu d'entraînement, "
    "téléchargées par scraping.")
text(s, M, Inches(2.4), Inches(6), Inches(1.3), "GANisme", size=80, bold=True, font=HEAD)
text(s, M, Inches(3.8), Inches(6), Inches(0.6), "Peindre avec un GAN", size=28, font=HEAD, color=GOLD)
text(s, M, Inches(5.3), Inches(6), Inches(0.7),
     ["La Plateforme · Deep Learning", "Équipe : [à compléter]"], size=14, color=MUTED)
side, gap = Inches(1.2), Inches(0.08)
for k, path in enumerate(safe(20, 1)):
    r, c = divmod(k, 4)
    picture(s, path, Inches(7.55) + c * (side + gap), Inches(0.55) + r * (side + gap), side)

# ================================================================ 2. Plan
s = slide_(
    "Le plan suit les attendus de l'énoncé : d'abord la veille scientifique sur les GAN, "
    "puis l'analyse du catalogue, la constitution du jeu d'images, le modèle, "
    "et enfin l'organisation de l'équipe et la conclusion.")
header(s, "Plan", "Du catalogue au générateur")
steps = ["Veille", "Catalogue", "Images", "Modèle", "Bilan"]
sw = Inches(2.2)
for k, word in enumerate(steps):
    x = M + k * (sw + Inches(0.27))
    card(s, x, Inches(2.1), sw, Inches(3.6))
    circle(s, x + (sw - Inches(0.9)) / 2, Inches(2.5), Inches(0.9), str(k + 1), size=24)
    text(s, x, Inches(3.75), sw, Inches(0.6), word, size=24, bold=True, font=HEAD, align=PP_ALIGN.CENTER)
    picture(s, safe(5, 40)[k], x + (sw - Inches(1.0)) / 2, Inches(4.45), Inches(1.0))

# ================================================================ Veille
divider("01", "Veille", ["15 questions", "GAN", "DCGAN"],
        "Première partie : la veille scientifique. Nous avons répondu aux 15 questions de l'énoncé ; "
        "les mots-clés sont à l'écran, les explications dans ces commentaires.", seed=2)

# Q1-Q2 : schema du GAN
s = slide_(
    "Q1 — Problème initial : le GAN (Goodfellow, 2014) répond à l'apprentissage NON SUPERVISÉ d'un "
    "modèle génératif : produire de nouvelles données qui ressemblent à un jeu d'exemples, sans étiquettes.\n"
    "Il fonctionne aussi pour l'image-à-image, la super-résolution, l'augmentation de données ou le "
    "semi-supervisé. Il est moins adapté au texte : les mots sont discrets, le gradient ne passe pas.\n\n"
    "Q2 — Fonctionnement (suivre le schéma) : le générateur G transforme un bruit aléatoire z en image. "
    "Le discriminateur D reçoit tantôt une vraie peinture, tantôt une image de G, et répond « vrai » ou "
    "« faux ». Le gradient de D indique à G comment mieux le tromper. Les deux progressent l'un contre "
    "l'autre (jeu à somme nulle) jusqu'à ce que les faux soient indiscernables.\n\n"
    "Ce que nos recherches montrent : le GAN n'est pas un classifieur. Le discriminateur classe, mais le but du GAN est de créer.")
header(s, "Veille · Q1–Q2", "Deux réseaux en compétition")
chip(s, M, Inches(1.75), Inches(2.6), "Non supervisé", fill=RED)
chip(s, M + Inches(2.8), Inches(1.75), Inches(2.6), "Jeu à somme nulle", fill=RED)
top, low, img = Inches(2.55), Inches(4.2), Inches(1.25)
bw, bh = Inches(1.9), Inches(1.0)
mid = low + (img - bh) / 2
node(s, M, mid, bw, bh, "Bruit z")
node(s, M + Inches(2.5), mid, bw, bh, "Générateur G", fill=RED)
arrow(s, M + bw + Inches(0.12), low + Inches(0.48))
arrow(s, M + Inches(2.5) + bw + Inches(0.12), low + Inches(0.48))
x_img = M + Inches(5.0)
picture(s, image_path(1027), x_img, top, img)
label(s, x_img - Inches(0.3), top + img, img + Inches(0.6), "vraie", size=14, color=MUTED, bold=False)
node(s, x_img, low, img, img, "?", size=44, color=GOLD)
label(s, x_img - Inches(0.3), low + img, img + Inches(0.6), "générée", size=14, color=MUTED, bold=False)
for row in (top, low):
    arrow(s, x_img + img + Inches(0.12), row + Inches(0.48))
x_d = x_img + img + Inches(0.6)
node(s, x_d, top, Inches(2.2), low + img - top, "Discriminateur D", sub="vrai / faux ?", size=20)
arrow(s, x_d + Inches(2.32), Inches(3.7))
text(s, x_d + Inches(2.8), Inches(3.35), Inches(2.3), Inches(1.0), ["gradient", "→ G apprend"],
     size=18, bold=True, color=GOLD, anchor=MSO_ANCHOR.MIDDLE)
trap(s, "Le GAN sert à classer", "Il sert à créer de nouvelles images")

# Q3 : discriminant vs generatif (nuages de points)
s = slide_(
    "Q3 — Rôles : le générateur crée des images ; le discriminateur est un classifieur binaire vrai/faux "
    "qui sert de critique au générateur.\n\n"
    "Différence entre modèles (suivre les deux schémas) :\n"
    "• À gauche, un modèle DISCRIMINANT apprend p(y|x) : il trace une frontière entre deux classes "
    "(par exemple portrait ou paysage). Il ne sait pas fabriquer de nouvel exemple.\n"
    "• À droite, un modèle GÉNÉRATIF apprend p(x) : la forme du nuage de données lui-même. Il peut donc "
    "produire de nouveaux points (en doré) qui ressemblent aux données.\n\n"
    "Ce que nos recherches montrent : dans un GAN, le discriminateur EST un modèle discriminant. Mais le GAN dans son ensemble est "
    "un modèle génératif : D n'est qu'un outil d'entraînement, c'est G qu'on garde à la fin.")
header(s, "Veille · Q3", "Discriminant ou génératif ?")
pw, ph, py = Inches(5.9), Inches(4.0), Inches(1.8)
rng = random.Random(3)
cluster_a = [(1.5 + rng.gauss(0, 0.45), 2.2 + rng.gauss(0, 0.45)) for _ in range(14)]
cluster_b = [(4.3 + rng.gauss(0, 0.45), 1.9 + rng.gauss(0, 0.45)) for _ in range(14)]
for k, (title, sub) in enumerate((("Discriminant", "apprend p(y | x)"), ("Génératif", "apprend p(x)"))):
    px = M + k * (pw + Inches(0.33))
    card(s, px, py, pw, ph)
    label(s, px + Inches(0.3), py + Inches(0.2), Inches(2.6), title, size=22, color=GOLD, align=PP_ALIGN.LEFT)
    label(s, px + Inches(3.0), py + Inches(0.2), Inches(2.6), sub, size=16, color=MUTED, bold=False,
          align=PP_ALIGN.RIGHT)
    for (cx, cy), color in [(p, RED) for p in cluster_a] + [(p, MUTED) for p in cluster_b]:
        dot(s, px + Inches(cx), py + Inches(0.6) + Inches(cy), Inches(0.16), color if k == 0 else STONE)
    if k == 0:
        line(s, px + Inches(3.25), py + Inches(0.75), px + Inches(2.65), py + ph - Inches(0.2), GOLD, 3, dash=True)
        label(s, px + Inches(3.0), py + ph - Inches(0.6), Inches(2.7), "frontière", size=15, color=GOLD)
    else:
        for cx, cy in [(1.2, 1.7), (2.0, 2.6), (1.6, 3.0), (4.0, 1.4), (4.7, 2.4), (4.4, 1.6)]:
            dot(s, px + Inches(cx), py + Inches(0.6) + Inches(cy), Inches(0.24), None, outline=GOLD)
        label(s, px + Inches(3.0), py + ph - Inches(0.6), Inches(2.7), "○ nouveaux exemples", size=15, color=GOLD)
trap(s, "Le GAN est un modèle discriminant", "D l'est, mais le GAN est génératif")

# Q4-Q5 : entrees, densite implicite
s = slide_(
    "Q4 — Entrées (schéma de gauche) : le générateur reçoit un vecteur de bruit z, par exemple 100 nombres "
    "tirés d'une loi normale (plus une étiquette dans un GAN conditionnel). Le discriminateur reçoit une "
    "image, vraie ou générée, et renvoie une probabilité : 0,87 veut dire « vraie à 87 % ».\n\n"
    "Q5 — Densité implicite (schéma de droite) : un modèle EXPLICITE (VAE, flows) sait calculer p(x), la "
    "probabilité d'une image donnée. Un GAN ne le sait pas : il sait seulement TIRER de nouvelles images, "
    "en passant du bruit dans G. La distribution n'existe qu'à travers cet échantillonnage : on dit "
    "qu'elle est implicite.\n\n"
    "Ce que nos recherches montrent : le GAN apprend bien la distribution des données, mais il ne peut pas dire quelle est la "
    "probabilité d'une image.")
header(s, "Veille · Q4–Q5", "Ce qui entre, ce qui sort")
card(s, M, Inches(1.8), Inches(5.6), Inches(4.0))
label(s, M + Inches(0.3), Inches(1.95), Inches(5), "Q4 · Entrées", size=22, color=GOLD, align=PP_ALIGN.LEFT)
# ligne 1 : z -> G -> image ; ligne 2 : image -> D -> probabilite
for j in range(6):
    card(s, M + Inches(0.45), Inches(2.65) + j * Inches(0.2), Inches(0.42), Inches(0.16), fill=STONE)
label(s, M + Inches(0.1), Inches(3.85), Inches(1.1), "z (100)", size=13, color=MUTED, bold=False)
arrow(s, M + Inches(1.0), Inches(3.05))
node(s, M + Inches(1.45), Inches(2.8), Inches(1.2), Inches(0.8), "G", fill=RED, size=22)
arrow(s, M + Inches(2.75), Inches(3.05))
node(s, M + Inches(3.2), Inches(2.8), Inches(1.6), Inches(0.8), "image", size=16, color=MUTED)
picture(s, image_path(696), M + Inches(0.4), Inches(4.4), Inches(0.95))
label(s, M + Inches(0.1), Inches(5.35), Inches(1.6), "vraie ou générée", size=13, color=MUTED, bold=False)
arrow(s, M + Inches(1.5), Inches(4.75))
node(s, M + Inches(1.95), Inches(4.5), Inches(1.2), Inches(0.8), "D", fill=RED, size=22)
arrow(s, M + Inches(3.25), Inches(4.75))
node(s, M + Inches(3.7), Inches(4.5), Inches(1.5), Inches(0.8), "0,87", size=20, color=GOLD,
     sub="vraie à 87 %")
card(s, Inches(6.55), Inches(1.8), Inches(6.18), Inches(4.0))
label(s, Inches(6.85), Inches(1.95), Inches(5.5), "Q5 · Densité", size=22, color=GOLD, align=PP_ALIGN.LEFT)
label(s, Inches(6.85), Inches(2.65), Inches(1.6), "Explicite", size=16, color=MUTED, align=PP_ALIGN.LEFT)
node(s, Inches(8.45), Inches(2.55), Inches(0.75), Inches(0.65), "x")
arrow(s, Inches(9.3), Inches(2.72))
node(s, Inches(9.75), Inches(2.55), Inches(1.25), Inches(0.65), "VAE")
arrow(s, Inches(11.1), Inches(2.72))
label(s, Inches(11.4), Inches(2.6), Inches(1.25), "p(x) ✓", size=17, color=TEXT)
label(s, Inches(6.85), Inches(4.05), Inches(1.6), "Implicite", size=16, color=GOLD, align=PP_ALIGN.LEFT)
node(s, Inches(8.45), Inches(3.95), Inches(0.75), Inches(0.65), "z")
arrow(s, Inches(9.3), Inches(4.12))
node(s, Inches(9.75), Inches(3.95), Inches(1.25), Inches(0.65), "G", fill=RED)
arrow(s, Inches(11.1), Inches(4.12))
label(s, Inches(11.4), Inches(4.0), Inches(1.25), "x ✓", size=17, color=GOLD)
label(s, Inches(8.45), Inches(4.85), Inches(4.2), "p(x) = ?   impossible à calculer", size=16, color=RED)
trap(s, "Le GAN calcule p(x)", "Il tire des images, sans formule de p(x)")

# Q6 : retropropagation alternee
s = slide_(
    "Q6 — Rétropropagation : l'entraînement alterne deux pas à chaque lot d'images (suivre les deux schémas).\n"
    "Pas 1, le discriminateur apprend : on lui montre des vraies images (étiquette 1) et des fausses "
    "(étiquette 0). On calcule l'entropie croisée binaire et on met à jour SEULEMENT D. G est figé (gris).\n"
    "Pas 2, le générateur apprend : G produit des images, D les juge. La perte de G est −log D(G(z)). "
    "Le gradient (flèche dorée) traverse D sans le modifier, jusqu'à G. On met à jour SEULEMENT G.\n\n"
    "Ce que nos recherches montrent : on n'entraîne jamais G et D en même temps sur la même perte. À chaque pas, l'un des deux est "
    "figé, sinon D apprendrait à se laisser tromper.")
header(s, "Veille · Q6", "Un pas pour D, un pas pour G")
for k, (title, g_fill, d_fill) in enumerate((("1 · D apprend", STONE, RED), ("2 · G apprend", RED, STONE))):
    px = M + k * (Inches(6.23))
    card(s, px, Inches(1.8), Inches(5.9), Inches(4.0))
    label(s, px + Inches(0.3), Inches(1.95), Inches(5.3), title, size=22, color=GOLD, align=PP_ALIGN.LEFT)
    node(s, px + Inches(0.35), Inches(3.05), Inches(1.4), Inches(1.0), "G", fill=g_fill, size=24,
         sub="figé" if k == 0 else "mis à jour")
    arrow(s, px + Inches(1.85), Inches(3.4))
    node(s, px + Inches(2.3), Inches(3.05), Inches(1.4), Inches(1.0), "D", fill=d_fill, size=24,
         sub="mis à jour" if k == 0 else "figé")
    arrow(s, px + Inches(3.8), Inches(3.4))
    node(s, px + Inches(4.25), Inches(3.05), Inches(1.35), Inches(1.0), "perte", size=18)
    line(s, px + Inches(4.9), Inches(4.25), px + Inches(4.9), Inches(4.75), GOLD, 3)
    target = px + Inches(3.0) if k == 0 else px + Inches(1.05)
    line(s, px + Inches(4.9), Inches(4.75), target, Inches(4.75), GOLD, 3, dash=k == 1)
    line(s, target, Inches(4.75), target, Inches(4.15), GOLD, 3, head=True)
    label(s, px + Inches(0.3), Inches(4.9), Inches(5.3), "gradient" if k == 0 else "gradient à travers D",
          size=15, color=GOLD, bold=False)
    if k == 0:
        label(s, px + Inches(2.0), Inches(2.5), Inches(2.0), "vraies = 1 · fausses = 0", size=13, color=MUTED, bold=False)
    else:
        label(s, px + Inches(2.0), Inches(2.5), Inches(2.0), "−log D(G(z))", size=14, color=MUTED, bold=False)
trap(s, "G et D apprennent ensemble", "On alterne : l'autre réseau est figé")

# Q7 : precision / rappel / FID
s = slide_(
    "Q7 — Évaluer un GAN est difficile : aucune perte ne mesure directement la qualité. Le schéma montre "
    "deux nuages : les vraies images (gris) et les images générées (doré).\n"
    "• PRÉCISION = la part des images générées qui tombent dans la zone des vraies : sont-elles réalistes ?\n"
    "• RAPPEL = la part des vraies images couverte par les générées : toute la variété est-elle produite ? "
    "Un mode collapse donne une bonne précision mais un rappel très faible.\n"
    "• FID (Fréchet Inception Distance) = la distance entre les deux nuages, mesurée sur des "
    "caractéristiques extraites par le réseau Inception. Plus bas = mieux. KID : même idée, plus fiable "
    "sur petits échantillons.\n"
    "• Et toujours : inspection visuelle, et recherche des plus proches voisins pour vérifier que G ne "
    "recopie pas les images d'entraînement.\n\n"
    "Ce que nos recherches montrent : une perte basse ne veut pas dire de belles images. On croise plusieurs métriques et l'œil.")
header(s, "Veille · Q7", "Fidélité et diversité")
card(s, M, Inches(1.8), Inches(7.4), Inches(4.0))
o1 = dot(s, M + Inches(2.9), Inches(3.8), Inches(3.4), None, outline=MUTED)
o2 = dot(s, M + Inches(4.3), Inches(3.8), Inches(2.4), None, outline=GOLD)
label(s, M + Inches(0.6), Inches(1.95), Inches(2.4), "vraies", size=17, color=MUTED)
label(s, M + Inches(4.5), Inches(1.95), Inches(2.4), "générées", size=17, color=GOLD)
line(s, M + Inches(2.9), Inches(3.8), M + Inches(4.3), Inches(3.8), RED, 3, head=True)
label(s, M + Inches(2.75), Inches(3.85), Inches(1.7), "FID", size=17, color=RED)
for j, (name, desc, fill) in enumerate((("Précision", "réalisme", RED), ("Rappel", "variété", CARD2),
                                        ("FID · KID", "distance", CARD2), ("Œil", "artefacts", CARD2),
                                        ("Voisins", "mémorisation", CARD2))):
    y = Inches(1.8) + j * Inches(0.82)
    chip(s, Inches(8.3), y, Inches(2.0), name, fill=fill, size=16, h=Inches(0.65))
    label(s, Inches(10.45), y + Inches(0.1), Inches(2.3), desc, size=16, color=MUTED, bold=False, align=PP_ALIGN.LEFT)
trap(s, "Perte basse = bonnes images", "Croiser métriques et inspection visuelle")

# Q8 : taille des jeux de donnees (echelle log)
s = slide_(
    "Q8 — Quantité de données : il n'existe pas de seuil universel ; cela dépend de la résolution visée et "
    "de la diversité des images. Le graphique, en échelle logarithmique, situe notre projet :\n"
    "• MetFaces : 1 336 portraits issus de musées, suffisants pour StyleGAN2-ADA grâce à l'augmentation "
    "adaptative (Karras et al., 2020).\n"
    "• Notre projet : environ 8 000 peintures, en 64×64.\n"
    "• CelebA : environ 200 000 visages, jeu classique pour le DCGAN.\n"
    "• LSUN chambres : environ 3 millions d'images.\n"
    "Avec trop peu de données, le discriminateur sur-apprend et le générateur finit par recopier les "
    "images d'entraînement.\n\n"
    "Ce que nos recherches montrent : « il faut des millions d'images » est faux. Mais plus les images sont variées (et c'est notre "
    "cas : 10 genres, 7 siècles), plus il en faut.")
header(s, "Veille · Q8", "Combien d'images ?")
data = CategoryChartData()
data.categories = ["MetFaces (ADA)", "Notre projet", "CelebA", "LSUN chambres"]
data.add_series("Images", [1336, 8000, 202599, 3033042])
chart = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, M, Inches(1.7), Inches(8.6), Inches(4.2), data).chart
dark_chart(chart, size=14)
chart.has_legend = False
plot = chart.plots[0]
plot.gap_width = 60
plot.has_data_labels = True
plot.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
series = plot.series[0]
for point, value in zip(series.points, [1336, 8000, 202599, 3033042]):
    tf = point.data_label.text_frame
    tf.text = fmt(value)
    run = tf.paragraphs[0].runs[0]
    run.font.size, run.font.bold, run.font.color.rgb = Pt(15), True, TEXT
    point.data_label.position = XL_LABEL_POSITION.OUTSIDE_END
series.format.fill.solid()
series.format.fill.fore_color.rgb = STONE
series.points[1].format.fill.solid()
series.points[1].format.fill.fore_color.rgb = GOLD
scaling = chart.value_axis._element.find(qn("c:scaling"))
scaling.insert(0, scaling.makeelement(qn("c:logBase"), {"val": "10"}))
chart.value_axis.visible = False
chart.value_axis.has_major_gridlines = False
chart.category_axis.tick_labels.font.color.rgb = TEXT
chart.category_axis.format.line.color.rgb = STONE
big_number(s, Inches(9.4), Inches(2.4), Inches(3.3), "~8 000", "nos peintures", size=48)
chip(s, Inches(9.6), Inches(4.4), Inches(2.9), "échelle log", fill=CARD2, size=14)
trap(s, "Il faut des millions d'images", "Ça dépend de la résolution et de la diversité")

# Q9 : convergence, mode collapse
s = slide_(
    "Q9 — Convergence : on ne cherche pas le minimum d'une perte, mais un équilibre entre G et D.\n"
    "• À gauche : si D devient trop fort, il rejette tout ; D(G(z)) ≈ 0 et le gradient reçu par G "
    "s'évanouit : G n'apprend plus.\n"
    "• Au centre : l'équilibre théorique. Les faux suivent exactement la distribution réelle et D répond "
    "0,5 partout, comme une pièce de monnaie.\n"
    "• À droite : le mode collapse. Des bruits z différents donnent tous la même image, parce qu'elle "
    "trompe D. G a « gagné » sans rien apprendre de la variété des peintures.\n"
    "En pratique, les pertes de G et D oscillent et ne descendent pas régulièrement.\n\n"
    "Ce que nos recherches montrent : une perte de G qui baisse ne veut pas dire que G progresse. Il faut regarder les images "
    "générées à intervalles réguliers.")
header(s, "Veille · Q9", "Un équilibre, pas un minimum")
pw = Inches(3.85)
for k, (title, fill) in enumerate((("D trop fort", CARD), ("Équilibre", CARD), ("Mode collapse", CARD))):
    px = M + k * (pw + Inches(0.29))
    card(s, px, Inches(1.8), pw, Inches(4.0), fill=fill)
    label(s, px, Inches(1.95), pw, title, size=22, color=RED if k != 1 else GOLD)
# D trop fort : grand D, petit G, gradient nul
node(s, M + Inches(0.4), Inches(3.3), Inches(0.9), Inches(0.75), "G", fill=STONE, size=18)
node(s, M + Inches(2.0), Inches(2.85), Inches(1.5), Inches(1.65), "D", fill=RED, size=30)
line(s, M + Inches(2.0), Inches(3.68), M + Inches(1.3), Inches(3.68), GOLD, 1, head=True, dash=True)
label(s, M, Inches(4.85), pw, "gradient ≈ 0", size=18, color=GOLD)
# Equilibre
px = M + pw + Inches(0.29)
big_number(s, px, Inches(2.9), pw, "D = 0,5", "partout", size=44)
# Mode collapse : plusieurs z vers une seule image
px = M + 2 * (pw + Inches(0.29))
for j in range(5):
    y = Inches(2.75) + j * Inches(0.5)
    dot(s, px + Inches(0.6), y + Inches(0.12), Inches(0.24), MUTED)
    line(s, px + Inches(0.75), y + Inches(0.12), px + Inches(2.35), Inches(3.88), STONE, 1.75, head=True)
picture(s, image_path(188), px + Inches(2.45), Inches(3.3), Inches(1.15))
label(s, px, Inches(4.85), pw, "5 bruits → 1 image", size=18, color=GOLD)
trap(s, "La perte de G baisse : il progresse", "Les pertes oscillent : regarder les images")

# Q10-Q11 : frise des variantes
s = slide_(
    "Q10 — Avantages : images nettes et réalistes, génération très rapide (une seule passe dans G), "
    "pas de vraisemblance à calculer, apprentissage non supervisé. Inconvénients : entraînement instable "
    "et sensible aux réglages, mode collapse, évaluation difficile.\n\n"
    "Q11 — Variantes, sur la frise :\n"
    "• 2014 : GAN (Goodfellow) et GAN conditionnel, qui ajoute une étiquette pour choisir la classe "
    "générée (par exemple un genre de peinture).\n"
    "• 2015 : DCGAN, réseaux convolutifs pour les images (notre baseline).\n"
    "• 2016 : Pix2Pix, traduction d'image à image avec des paires (croquis → photo) ; SRGAN, "
    "super-résolution.\n"
    "• 2017 : CycleGAN, traduction sans paires (photo → tableau de Monet) ; WGAN-GP, entraînement plus "
    "stable grâce à la distance de Wasserstein.\n"
    "• 2018 : StyleGAN, visages ultra-réalistes (site thispersondoesnotexist).")
header(s, "Veille · Q10–Q11", "Dix ans de variantes")
y0 = Inches(3.6)
line(s, M, y0, W - M, y0, STONE, 3)
events = [("2014", "GAN · cGAN", "par classe"), ("2015", "DCGAN", "images"), ("2016", "Pix2Pix · SRGAN", "croquis → photo"),
          ("2017", "CycleGAN · WGAN-GP", "photo → Monet"), ("2018", "StyleGAN", "visages")]
step = (W - 2 * M) // len(events)
for k, (year, names, use) in enumerate(events):
    cx = M + step * k + step // 2
    dot(s, cx, y0, Inches(0.4), GOLD if k == 1 else RED)
    label(s, cx - step // 2, Inches(2.75), step, year, size=24, color=GOLD, h=Inches(0.5))
    label(s, cx - step // 2, Inches(4.0), step, names, size=17, h=Inches(0.5))
    label(s, cx - step // 2, Inches(4.5), step, use, size=15, color=MUTED, bold=False, h=Inches(0.45))
chip(s, M, Inches(1.85), Inches(3.4), "+ nettes · rapides", fill=CARD2, size=16)
chip(s, M + Inches(3.6), Inches(1.85), Inches(3.4), "− instables · collapse", fill=RED, size=16)
chip(s, M, Inches(5.3), Inches(3.4), "DCGAN = notre baseline", fill=CARD2, size=15, color=GOLD)

# Q12 : CNN + GAN = DCGAN
s = slide_(
    "Q12 — Un CNN (réseau convolutif) est une ARCHITECTURE de réseau, en général utilisée pour classer "
    "des images. Un GAN est une MÉTHODE D'ENTRAÎNEMENT à deux réseaux, quelle que soit leur architecture.\n"
    "Le vanilla GAN de 2014 utilisait des couches denses : il marchait sur des chiffres manuscrits mais "
    "restait instable sur des images complexes.\n"
    "Le DCGAN (Radford et al., 2015) est un GAN dont G et D sont des CNN, avec des règles précises : "
    "convolutions à pas au lieu du pooling, convolutions transposées dans G pour agrandir l'image, "
    "BatchNorm, ReLU dans G et LeakyReLU dans D, tanh en sortie. Résultat : un entraînement beaucoup "
    "plus stable sur des images, jusqu'à environ 64×64.\n\n"
    "Ce que nos recherches montrent : la question « comparez un CNN et un GAN » compare deux choses de nature différente. "
    "Il faut le dire : l'un est une architecture, l'autre une méthode ; le DCGAN combine les deux.")
header(s, "Veille · Q12", "CNN, GAN, DCGAN")
trio = [("CNN", "architecture", "classer"), ("GAN", "méthode", "générer"),
        ("DCGAN", "GAN + convolutions", "images 64×64")]
tw = Inches(3.6)
for j, (name, what, use) in enumerate(trio):
    x = M + j * (tw + Inches(0.65))
    card(s, x, Inches(1.8), tw, Inches(3.95), fill=RED if j == 2 else CARD)
    text(s, x, Inches(2.15), tw, Inches(0.9), name, size=40, bold=True, font=HEAD, align=PP_ALIGN.CENTER,
         color=TEXT if j == 2 else GOLD)
    chip(s, x + Inches(0.4), Inches(3.45), tw - Inches(0.8), what, size=17, fill=BG if j == 2 else CARD2)
    chip(s, x + Inches(0.4), Inches(4.3), tw - Inches(0.8), use, size=17, fill=BG if j == 2 else CARD2)
    if j < 2:
        label(s, x + tw + Inches(0.1), Inches(3.4), Inches(0.45), "+" if j == 0 else "=", size=40, color=MUTED,
              h=Inches(0.8))
trap(s, "Un GAN est un type de CNN", "Architecture ≠ méthode : le DCGAN combine")

# Q13 : arbre des modeles generatifs
s = slide_(
    "Q13 — Le chatbot dont parle l'énoncé est ChatGPT. Il repose sur GPT (Generative Pre-trained "
    "Transformer), de la famille des TRANSFORMERS : de grands modèles de langage autorégressifs, qui "
    "prédisent le mot suivant, affinés grâce aux retours humains (RLHF). Il a été développé par OpenAI.\n"
    "L'arbre montre les grandes familles de modèles génératifs profonds : les GAN (notre projet), les "
    "VAE, les modèles de diffusion (utilisés par DALL-E 2) et les transformers (GPT, ChatGPT).\n\n"
    "Ce que nos recherches montrent : ChatGPT n'est pas un GAN, et DALL-E non plus. Ce sont deux autres familles de modèles "
    "génératifs.")
header(s, "Veille · Q13", "Les familles de modèles génératifs")
node(s, Inches(4.67), Inches(1.75), Inches(4.0), Inches(0.8), "Modèles génératifs", fill=CARD2, size=20)
fams = [("GAN", "notre projet", RED), ("VAE", "densité explicite", CARD2), ("Diffusion", "DALL-E 2", CARD2),
        ("Transformer", "GPT → ChatGPT", GOLD)]
fw = Inches(2.7)
for k, (name, sub, fill) in enumerate(fams):
    x = M + k * (fw + Inches(0.41))
    line(s, Inches(6.67), Inches(2.55), x + fw / 2, Inches(3.3), STONE, 2)
    node(s, x, Inches(3.3), fw, Inches(1.25), name, fill=fill, size=22, color=BG if fill == GOLD else TEXT)
    label(s, x, Inches(4.65), fw, sub, size=16, color=GOLD if k == 3 else MUTED, bold=k == 3)
label(s, M + 3 * (fw + Inches(0.41)), Inches(5.15), fw, "OpenAI", size=15, color=MUTED, bold=False)
trap(s, "ChatGPT est un GAN", "C'est un transformer (GPT, OpenAI)")

# Q14-Q15
s = slide_(
    "Q14 — Dangers des modèles génératifs comme les GAN :\n"
    "• deepfakes : usurpation d'identité, fraude, chantage ;\n"
    "• désinformation à grande échelle (fausses images crédibles) ;\n"
    "• biais : le modèle reproduit les déséquilibres de ses données (notre catalogue est très européen et "
    "très religieux, par exemple) ;\n"
    "• droits d'auteur : imitation du style d'artistes vivants ;\n"
    "• vie privée : mémorisation possible d'images d'entraînement.\n\n"
    "Q15 — Intérêt de générer des données : compléter des données rares (imagerie médicale), produire des "
    "données synthétiques anonymes, offrir un outil de création aux artistes, explorer des styles.\n"
    "Notre position : c'est utile à condition que le contenu synthétique reste identifié comme tel.")
header(s, "Veille · Q14–Q15", "Risques et intérêt")
risks = [("Deepfakes", RED), ("Désinformation", RED), ("Biais", CARD2), ("Droits d'auteur", CARD2),
         ("Vie privée", CARD2)]
label(s, M, Inches(1.8), Inches(6), "Q14 · Dangers", size=22, color=GOLD, align=PP_ALIGN.LEFT)
for j, (word, fill) in enumerate(risks):
    r, c = divmod(j, 2)
    chip(s, M + c * Inches(3.0), Inches(2.5) + r * Inches(0.95), Inches(2.8), word, fill=fill, size=18, h=Inches(0.7))
card(s, Inches(7.0), Inches(1.8), Inches(5.73), Inches(4.0))
label(s, Inches(7.3), Inches(1.95), Inches(5), "Q15 · Pourquoi générer ?", size=22, color=GOLD, align=PP_ALIGN.LEFT)
for j, word in enumerate(["Données rares", "Anonymisation", "Création artistique"]):
    chip(s, Inches(7.3), Inches(2.75) + j * Inches(0.85), Inches(5.1), word, fill=CARD2, size=18, h=Inches(0.65))
label(s, Inches(7.3), Inches(5.25), Inches(5.1), "si le synthétique est signalé", size=16, color=GOLD, bold=False)

# Synthese de la veille
s = slide_(
    "Synthèse des enseignements de notre veille, à retenir pour l'oral :\n"
    "1. Un GAN ne sert pas à classer : il sert à créer. (Q1)\n"
    "2. Le discriminateur est discriminant, mais le GAN est génératif. (Q3)\n"
    "3. Un GAN ne calcule pas p(x) : sa densité est implicite. (Q5)\n"
    "4. G et D ne s'entraînent pas en même temps : on alterne. (Q6)\n"
    "5. Une perte basse ne prouve rien : on croise métriques et images. (Q7, Q9)\n"
    "6. CNN et GAN ne se comparent pas directement : architecture contre méthode. (Q12)\n"
    "7. ChatGPT et DALL-E ne sont pas des GAN. (Q13)")
header(s, "Veille · synthèse", "Ce que notre veille a montré")
findings = [("Q1", "GAN = créer, pas classer"), ("Q3", "D discriminant, GAN génératif"),
              ("Q5", "Pas de p(x) calculable"), ("Q6", "On alterne G et D"),
              ("Q7·9", "Perte basse ≠ belles images"), ("Q12", "Architecture ≠ méthode"),
              ("Q13", "ChatGPT n'est pas un GAN")]
for j, (q, txt) in enumerate(findings):
    r, c = divmod(j, 2)
    x, y = M + c * Inches(6.2), Inches(1.85) + r * Inches(1.2)
    card(s, x, y, Inches(5.93), Inches(1.0))
    circle(s, x + Inches(0.2), y + Inches(0.12), Inches(0.76), q, size=14 if len(q) > 3 else 16)
    label(s, x + Inches(1.15), y + Inches(0.27), Inches(4.6), txt, size=19, align=PP_ALIGN.LEFT)

# ================================================================ Catalogue
divider("02", "Catalogue", ["52 867 œuvres", "Métadonnées", "Sans images"],
        "Deuxième partie : l'analyse du catalogue. Nous l'avons faite uniquement sur les métadonnées "
        "(le fichier CSV), avant de télécharger la moindre image, pour décider quoi télécharger. "
        "Les graphiques qui suivent sont ceux produits par notre notebook 01_eda.", seed=3)

s = slide_(
    "Le catalogue de la Web Gallery of Art contient 52 867 œuvres : peintures, sculptures, architecture, "
    "gravures… L'énoncé demande de ne garder que les peintures : 32 438, soit 61 %.\n"
    "Elles viennent de 3 663 artistes. 88 % des peintures ont une date connue ; les dates inconnues sont "
    "codées par un tiret, pas laissées vides.\n"
    "Chaque œuvre a une URL unique : c'est notre identifiant pour relier chaque image à sa source.\n"
    "Les champs sont du texte libre (« c. 1530 », « Oil on canvas, 68 x 95 cm ») : nous avons écrit "
    "des fonctions testées pour en extraire l'année, les dimensions et la technique.")
header(s, "Catalogue", "Les chiffres clés")
nums = [(fmt(EDA["n_catalogue"]), "œuvres"), (fmt(n_paint), "peintures"),
        (fmt(EDA["n_artists"]), "artistes"), (pct(EDA["share_known_year"]), "datées")]
sw = Inches(2.85)
for k, (value, label_) in enumerate(nums):
    r, c = divmod(k, 2)
    x, y = M + c * Inches(2.95), Inches(1.85) + r * Inches(1.95)
    card(s, x, y, Inches(2.8), Inches(1.75))
    big_number(s, x, y + Inches(0.2), Inches(2.8), value, label_, size=40)
figure(s, FIG_EDA / "01_forms.png", Inches(6.6), Inches(1.75), Inches(6.1), Inches(4.9), "Formes d'art du catalogue · notebook 01_eda")
for k, word in enumerate(["URL unique", "Texte libre", "Inconnu = « - »"]):
    chip(s, M + k * Inches(1.97), Inches(5.95), Inches(1.85), word, size=13)

s = slide_(
    f"Les sujets religieux représentent {pct(types['religious'] / n_paint)} des peintures, presque trois fois "
    "plus que les portraits, deuxième genre. Les natures mortes, intérieurs et scènes de genre sont "
    "minoritaires.\n"
    "Conséquence : si on tirait les images au hasard, le GAN produirait surtout des madones et des scènes "
    "bibliques. Décision : imposer des quotas équilibrés par genre lors de la sélection.")
header(s, "Catalogue", "Déséquilibre des genres")
figure(s, FIG_EDA / "02_types.png", M, Inches(1.75), Inches(8.3), Inches(5.0), "Graphique du notebook 01_eda")
big_number(s, Inches(9.2), Inches(2.4), Inches(3.5), pct(types["religious"] / n_paint), "religieux", color=RED, size=60)
chip(s, Inches(9.5), Inches(4.6), Inches(2.9), "→ quotas par genre", fill=CARD2, size=16)

s = slide_(
    "À gauche, les écoles : l'école italienne représente 42 % des peintures ; avec les écoles française, "
    "hollandaise et flamande, on dépasse 80 %. Le corpus est très européen de l'Ouest.\n"
    "À droite, les artistes les plus présents : Giotto arrive en tête, surtout parce que ses cycles de "
    "fresques sont détaillés vue par vue.\n"
    "Mais la concentration reste modérée : les 10 premiers artistes ne font que 10 % des peintures, et la "
    "moitié des artistes ont 3 peintures ou moins.\n"
    "Décision : un plafond de 30 peintures par artiste, comme garde-fou contre la mémorisation.")
header(s, "Catalogue", "Écoles et artistes")
figure(s, FIG_EDA / "03_schools_artists.png", M, Inches(1.75), Inches(12.1), Inches(4.6), "Graphiques du notebook 01_eda")
chip(s, M, Inches(6.55), Inches(3.4), "Italie : 42 %", fill=RED, size=16)
chip(s, M + Inches(3.6), Inches(6.55), Inches(3.8), "→ 30 peintures max / artiste", fill=CARD2, size=16)

s = slide_(
    "Répartition des dates de création. Le catalogue couvre surtout 1300-1900, avec trois pics : la "
    "Renaissance (vers 1500), le XVIIe siècle (baroque, âge d'or hollandais) et les années 1880 "
    "(impressionnisme, Van Gogh). L'art du XXe siècle est absent : le catalogue s'arrête vers 1900.\n"
    "Le graphique de droite montre que le genre dépend de l'époque : le religieux domine jusqu'au XVIe "
    "siècle, puis portraits, paysages et scènes de genre prennent le relais.\n"
    "Conséquence : équilibrer les genres équilibre aussi, en partie, les époques.\n"
    "Remarque : la colonne TIMEFRAME ne correspond pas à la date de l'œuvre dans 18 % des cas ; nous "
    "avons recalculé l'année à partir de DATE.")
header(s, "Catalogue", "Le genre suit l'époque")
figure(s, FIG_EDA / "05_years.png", M, Inches(1.75), Inches(5.9), Inches(4.4), "Années de création · 01_eda")
figure(s, FIG_EDA / "07_type_by_century.png", Inches(6.75), Inches(1.75), Inches(5.95), Inches(4.4), "Genre par siècle · 01_eda")
chip(s, M, Inches(6.55), Inches(3.6), "3 pics : 1500 · 1650 · 1880", fill=CARD2, size=15)
chip(s, Inches(6.75), Inches(6.55), Inches(3.8), "religieux : 99 % → 3 %", fill=RED, size=15)

s = slide_(
    "À gauche, les techniques : l'huile domine (70 %) ; les fresques (14,5 %) sont des photos de murs et "
    "de voûtes, que nous excluons.\n"
    "À droite, le ratio largeur/hauteur calculé à partir des dimensions en centimètres. Deux bosses : les "
    "formats verticaux (portraits, vers 0,8) et horizontaux (paysages, vers 1,33). Or un DCGAN génère des "
    "images carrées : un recadrage central fait perdre 24 % de la surface en médiane.\n"
    "Décision : recadrage carré central, et exclusion des formats trop allongés (ratio hors de 0,5 à 2), "
    "comme les frises ou les prédelles.")
header(s, "Catalogue", "Techniques et formats")
figure(s, FIG_EDA / "06_techniques_ratios.png", M, Inches(1.75), Inches(12.1), Inches(4.6), "Graphiques du notebook 01_eda")
chip(s, M, Inches(6.55), Inches(3.4), "Fresques exclues", fill=RED, size=16)
chip(s, M + Inches(3.6), Inches(6.55), Inches(3.8), "→ carré central", fill=CARD2, size=16)

s = slide_(
    "Quatre découvertes faites pendant l'analyse, chacune avec sa décision :\n"
    f"• {pct(EDA['n_detail_views'] / n_paint)} de vues de détail : ce sont des recadrages d'une autre œuvre du "
    "catalogue. Elles créent des quasi-doublons, et une fuite si une œuvre est en entraînement et son "
    "détail en évaluation. → Exclues, et découpage entraînement/évaluation par œuvre.\n"
    f"• {pct(EDA['techniques']['fresco'] / n_paint)} de fresques : photos de murs et de voûtes. → Exclues.\n"
    "• 18 % de périodes incohérentes : la colonne TIMEFRAME ne colle pas à la date. → On utilise DATE.\n"
    "• Ratio médian de 0,85 : les tableaux sont souvent verticaux. → Recadrage carré.")
header(s, "Catalogue", "Quatre découvertes dans les données")
traps = [(pct(EDA["n_detail_views"] / n_paint), "détails", "exclus"),
         (pct(EDA["techniques"]["fresco"] / n_paint), "fresques", "exclues"),
         (pct(1 - EDA["share_year_inside_timeframe"]), "dates incohérentes", "DATE"),
         ("0,85", "ratio médian", "carré central")]
for k, (value, label_, decision) in enumerate(traps):
    x = M + k * (sw + Inches(0.23))
    card(s, x, Inches(1.95), sw, Inches(4.7))
    big_number(s, x, Inches(2.5), sw, value, label_, size=50)
    chip(s, x + Inches(0.35), Inches(5.0), sw - Inches(0.7), "→ " + decision, fill=RED, size=16)

# ================================================================ Images
divider("03", "Images", ["Sélection", "Scraping", "Contrôles"],
        "Troisième partie : comment nous passons du catalogue à un jeu d'images prêt pour l'entraînement. "
        "Les graphiques viennent de notre notebook 02_dataset.", seed=4)

s = slide_(
    "Sélection en entonnoir : 32 438 peintures, puis sans vues de détail, sans fresques, avec un format "
    "raisonnable : il reste 23 228 candidates. Nous en tirons 8 000 avec une graine fixe (reproductible) : "
    "une image par œuvre, 30 peintures maximum par artiste, et des quotas équilibrés par genre.\n"
    "Le graphique (notebook 02_dataset) compare le catalogue (gris) et notre sélection (bleu) : le "
    "religieux passe de 42 % à 14 %, l'école italienne de 42 % à 24 %.\n"
    "Les œuvres sont réparties 90 % entraînement / 10 % évaluation, par œuvre : jamais des deux côtés.")
header(s, "Images", "8 000 peintures équilibrées")
funnel = EDA["selection_funnel"]
steps = [funnel["peintures"], funnel["sans vues de détail"], funnel["+ sans fresques"],
         funnel["+ ratio connu dans [0.5, 2] ou inconnu"], len(MANIFEST)]
fw = Inches(2.1)
for k, value in enumerate(steps):
    x = M + k * (fw + Inches(0.4))
    chip(s, x, Inches(1.8), fw, fmt(value), fill=RED if k == len(steps) - 1 else CARD2, size=20, h=Inches(0.7))
    if k < len(steps) - 1:
        arrow(s, x + fw + Inches(0.05), Inches(2.0), Inches(0.3))
figure(s, FIG_DS / "01_selection_vs_catalogue.png", M, Inches(2.8), Inches(12.1), Inches(4.1),
       "Catalogue vs sélection · notebook 02_dataset")

s = slide_(
    "Le scraping récupère chaque image sur le site wga.hu. Un fichier manifeste suit chaque image avec "
    "son statut : en attente, traitée, en échec (retentée plus tard) ou écartée. On peut arrêter et "
    "reprendre sans rien perdre.\n"
    "À gauche, un extrait réel du journal de téléchargement : le débit et l'état sont enregistrés toutes "
    "les 25 images.\n"
    f"État au {stamp} : {fmt(n_ok)} images traitées, {fmt(n_fail)} en échec, {fmt(n_skip)} écartées, "
    f"{fmt(n_pend)} en attente.\n"
    "Difficulté rencontrée : le premier téléchargement s'est bloqué sans message d'erreur. Le serveur "
    "envoyait certaines images au compte-gouttes, et notre limite de temps ne portait que sur chaque "
    "lecture réseau. Correction : un délai total de 45 secondes par image, vérifié par un test.")
header(s, "Images", "Un scraping reprenable")
terminal(s, M, Inches(1.8), Inches(7.3), Inches(3.4), LOG_LINES, "data/processed/download.log")
progress = [(fmt(n_ok), "traitées", GOLD), (fmt(n_fail), "échecs", RED), (fmt(n_skip), "écartées", MUTED),
            (fmt(n_pend), "en attente", MUTED)]
for k, (value, label_, color) in enumerate(progress):
    r, c = divmod(k, 2)
    big_number(s, Inches(8.2) + c * Inches(2.3), Inches(1.85) + r * Inches(1.75), Inches(2.2), value, label_,
               color=color, size=36)
chip(s, M, Inches(5.65), Inches(5.2), "Difficulté : blocage serveur", fill=CARD2, size=16)
chip(s, M + Inches(5.5), Inches(5.65), Inches(4.6), "→ délai de 45 s par image", fill=RED, size=16)

s = slide_(
    "Extrait réel du manifeste (data/processed/manifest.csv) : une ligne par image sélectionnée, avec son "
    "identifiant, l'artiste, le titre, le genre, l'ensemble (entraînement ou évaluation), le statut et "
    "la taille d'origine. C'est ce fichier qui rend le téléchargement reprenable et qui relie chaque "
    "image à sa source.")
header(s, "Images", "Le manifeste, image par image")
rows = MANIFEST[MANIFEST["image_id"].isin(SAFE_ROWS)].set_index("image_id").loc[SAFE_ROWS].reset_index()
rh = Inches(0.78)
cols_x = [Inches(1.55), Inches(2.45), Inches(5.35), Inches(8.85), Inches(10.4), Inches(11.5)]
heads = ["id", "artiste", "titre", "genre", "ensemble", "statut"]
for x, h_ in zip(cols_x, heads):
    text(s, x, Inches(1.8), Inches(2.8), Inches(0.35), h_.upper(), size=12, bold=True, color=GOLD)
for r, row in rows.iterrows():
    y = Inches(2.25) + r * (rh + Inches(0.08))
    card(s, M, y, Inches(12.13), rh, fill=CARD)
    picture(s, image_path(int(row["image_id"])), M + Inches(0.08), y + Inches(0.06), rh - Inches(0.12))
    values = [str(row["image_id"]), row["AUTHOR"][:26], row["TITLE"][:34], GENRE_FR[row["TYPE"]],
              row["split"], row["status"]]
    for c, (x, v) in enumerate(zip(cols_x, values)):
        text(s, x, y, Inches(3.4) if c == 2 else Inches(2.8), rh, v, size=14,
             color=GOLD if c == 5 else TEXT, bold=c == 5, anchor=MSO_ANCHOR.MIDDLE)

s = slide_(
    f"Contrôle des images téléchargées (notebook 02_dataset, {fmt(CHECKS['processed'])} images au moment de "
    "l'exécution).\n"
    f"À gauche : le petit côté des images d'origine. Médiane : {CHECKS['median_short_side_px']:.0f} pixels, "
    "très au-dessus de nos 256 pixels d'archive et des 64 pixels d'entraînement : aucune perte de détail utile.\n"
    "À droite : le ratio largeur/hauteur réel des images, comparé à celui calculé depuis les dimensions en "
    f"cm du catalogue. Ils concordent à ±10 % pour {pct(CHECKS['share_ratio_within_10pct'])} des images : "
    "les métadonnées étaient fiables. Les écarts viennent d'images recadrées ou photographiées avec leur cadre.")
header(s, "Images", "Les images d'origine")
figure(s, FIG_DS / "02_original_sizes.png", M, Inches(1.75), Inches(12.1), Inches(4.6), "Graphiques du notebook 02_dataset")
chip(s, M, Inches(6.55), Inches(3.6), f"médiane : {CHECKS['median_short_side_px']:.0f} px", fill=CARD2, size=16)
chip(s, M + Inches(3.8), Inches(6.55), Inches(4.2), f"ratio conforme : {pct(CHECKS['share_ratio_within_10pct'])}",
     fill=CARD2, size=16)

s = slide_(
    "Question 5 de l'énoncé : ce que montrent les images, et les conséquences sur le prétraitement.\n"
    "• Tons sombres et bruns : vernis anciens, fonds noirs. On ne corrige pas les couleurs : c'est le "
    "style des œuvres que le GAN doit apprendre.\n"
    "• Cadres visibles sur certaines photos, ou mur autour de l'œuvre : du bruit, en faible proportion.\n"
    "• Le sujet est souvent centré : le recadrage carré central le conserve.\n"
    "Les images de cette diapositive sont de vraies images de notre jeu, après recadrage en 256×256.")
header(s, "Images · Q5", "Ce que montrent les images")
rows_g = ["portrait", "landscape", "still-life", "interior", "genre"]
side, gap = Inches(0.95), Inches(0.08)
for r, genre in enumerate(rows_g):
    y = Inches(1.75) + r * (side + gap)
    text(s, M, y, Inches(1.6), side, GENRE_FR[genre], size=14, bold=True, color=MUTED, anchor=MSO_ANCHOR.MIDDLE)
    ids = SAFE_IMAGES[genre][:]
    random.Random(r).shuffle(ids)
    for c, image_id in enumerate(ids[:6]):
        picture(s, image_path(image_id), M + Inches(1.6) + c * (side + gap), y, side)
for j, (word, fill) in enumerate([("Tons sombres", CARD2), ("Cadres visibles", CARD2), ("Sujet centré", CARD2),
                                  ("Haute résolution", CARD2)]):
    chip(s, Inches(9.0), Inches(2.2) + j * Inches(1.0), Inches(3.7), word, fill=fill, size=18, h=Inches(0.7))

s = slide_(
    "Couleurs mesurées sur toutes les images téléchargées (notebook 02_dataset).\n"
    "À gauche, la saturation moyenne : la ligne orange est notre seuil de détection des reproductions en "
    f"noir et blanc. Résultat : {CHECKS['grayscale']} image(s) seulement sous le seuil. Nous craignions un "
    "problème de photos noir et blanc ; les mesures montrent qu'il est négligeable.\n"
    f"À droite, la luminosité : médiane de {str(CHECKS['median_brightness']).replace('.', ',')} sur une échelle "
    "de 0 (noir) à 1 (blanc). Les peintures sont globalement sombres, ce qui confirme l'observation "
    "visuelle. Nous gardons ces couleurs telles quelles : c'est le style à apprendre.")
header(s, "Images", "Des peintures sombres, très peu de noir et blanc")
figure(s, FIG_DS / "05_color_stats.png", M, Inches(1.75), Inches(12.1), Inches(4.6), "Graphiques du notebook 02_dataset")
chip(s, M, Inches(6.55), Inches(3.6), f"noir et blanc : {CHECKS['grayscale']} image(s)", fill=CARD2, size=16)
chip(s, M + Inches(3.8), Inches(6.55), Inches(3.6), "→ tons sombres conservés", fill=RED, size=16)

s = slide_(
    "Un vrai bug trouvé grâce aux visualisations.\n"
    "Pour repérer les quasi-doublons, nous calculions une empreinte de chaque image (aHash 8×8). En "
    "regardant les paires détectées, nous avons vu qu'il s'agissait presque toujours de paysages "
    "DIFFÉRENTS : ciel clair en haut, sol sombre en bas, donc même empreinte.\n"
    "Le graphique mesure l'écart réel des pixels de chaque paire candidate : "
    f"{CHECKS['near_duplicate_candidates']} candidates, mais une seule sous le seuil de 10, à gauche : "
    "un vrai doublon (la même nature morte attribuée à deux artistes). Toutes les autres sont des images "
    "différentes.\n"
    "Sans correctif, l'export aurait supprimé des centaines de paysages et biaisé le jeu d'images. "
    "Correction : l'empreinte propose des candidats, la comparaison des pixels confirme. Un test "
    "automatique couvre ce cas.")
header(s, "Images", "Un bug révélé par les visuels")
figure(s, FIG_DS / "07_near_duplicate_check.png", M, Inches(1.75), Inches(7.6), Inches(4.5),
       "Contrôle des quasi-doublons · notebook 02_dataset")
big_number(s, Inches(8.5), Inches(1.95), Inches(4.2), fmt(CHECKS["near_duplicate_candidates"]), "paires candidates",
           color=MUTED, size=44)
big_number(s, Inches(8.5), Inches(3.75), Inches(4.2), fmt(CHECKS["near_duplicate_confirmed"]), "vrai doublon",
           color=GOLD, size=44)
chip(s, Inches(8.7), Inches(5.6), Inches(3.8), "→ confirmation par pixels", fill=RED, size=16)

s = slide_(
    "Les choix de prétraitement, chacun justifié par une observation :\n"
    "• Recadrage carré central : le DCGAN génère des images carrées.\n"
    "• 64×64 pixels : baseline rapide à entraîner ; 128×128 possible sans retélécharger, les images étant "
    "archivées en 256×256.\n"
    "• Pixels normalisés entre −1 et 1 : la plage de sortie de la fonction tanh du générateur.\n"
    "• Exclusion des rares images en noir et blanc.\n"
    "• Suppression des quasi-doublons confirmés : évite la mémorisation et la fuite vers l'évaluation.\n"
    "• Miroir horizontal aléatoire : double la variété ; pas de miroir vertical (un ciel en bas n'est pas "
    "plausible).")
header(s, "Images", "Prétraitement")
pre = [("Carré", "central"), ("64×64", "pixels"), ("[−1, 1]", "sortie tanh"),
       ("Sans N&B", "saturation"), ("Sans doublons", "pixels confirmés"), ("Miroir", "horizontal")]
tw = Inches(3.85)
for j, (big, small) in enumerate(pre):
    r, c = divmod(j, 3)
    x, y = M + c * (tw + Inches(0.3)), Inches(1.9) + r * Inches(2.45)
    card(s, x, y, tw, Inches(2.2), fill=RED if j == 1 else CARD)
    text(s, x, y + Inches(0.4), tw, Inches(0.9), big, size=38, bold=True, font=HEAD,
         color=TEXT if j == 1 else GOLD, align=PP_ALIGN.CENTER)
    text(s, x, y + Inches(1.35), tw, Inches(0.5), small, size=17, color=TEXT if j == 1 else MUTED,
         align=PP_ALIGN.CENTER)

s = slide_(
    "La qualité du code est vérifiée par des tests automatiques (pytest). Ci-contre, le résultat réel de "
    "leur exécution au moment de générer cette présentation.\n"
    "Ils couvrent l'extraction des métadonnées (dates, dimensions, techniques), la sélection (quotas, "
    "plafonds, découpage sans fuite), le téléchargement (reprise, doublons, délai total par image) et "
    "les contrôles d'images (empreintes, confirmation des doublons).\n"
    "Chaque bug rencontré (blocage du téléchargement, faux doublons) a donné lieu à un test.")
header(s, "Images", "Un code testé")
terminal(s, M, Inches(1.8), Inches(7.6), Inches(4.2), TEST_LINES, "python -m pytest")
for j, word in enumerate(["Métadonnées", "Sélection", "Téléchargement", "Contrôles d'images"]):
    chip(s, Inches(8.6), Inches(1.9) + j * Inches(0.95), Inches(4.1), word, fill=CARD2, size=18, h=Inches(0.7))
chip(s, Inches(8.6), Inches(5.75), Inches(4.1), "1 bug → 1 test", fill=RED, size=18, h=Inches(0.7))

# ================================================================ Modele (enonce Q6 a Q9)
divider("04", "Modèle", ["DCGAN", "64×64", "FID · KID"],
        "Quatrième partie : le modèle, en suivant les questions 6 à 9 de l'énoncé : architecture et pertes, "
        "type de GAN et hyperparamètres, évaluation, conclusion.", seed=5)

hp = TRAIN.get("hyperparams", {})
s = slide_(
    "Question 6 de l'énoncé — architecture et fonctions de perte.\n"
    "Le générateur part d'un vecteur de bruit de 100 nombres et l'agrandit étape par étape avec des "
    "convolutions transposées : 4×4, 8×8, 16×16, 32×32, jusqu'à une image 64×64 en couleurs. Chaque étape "
    "utilise BatchNorm et ReLU ; la sortie utilise tanh, d'où la normalisation des pixels entre −1 et 1.\n"
    "Le discriminateur fait le chemin inverse avec des convolutions à pas (stride 2) et LeakyReLU 0,2, et "
    "termine par une probabilité « vraie ».\n"
    "Fonctions de perte : entropie croisée binaire (BCE). Le discriminateur minimise "
    "−[log D(x) + log(1 − D(G(z)))] ; le générateur minimise −log D(G(z)), la version « non saturante » "
    "proposée par Goodfellow, qui donne des gradients utiles au début de l'entraînement quand D gagne "
    "facilement.\n"
    "Alternative envisagée : la perte de Wasserstein avec pénalité de gradient (WGAN-GP), plus stable mais "
    "plus coûteuse (5 pas de critique par pas de G).\n"
    "En 128×128, on ajoute une couche dans chaque réseau : le générateur part de 4×4×1024 et passe par 64×64 "
    "avant d'atteindre 128×128 ; le discriminateur fait l'inverse. Le générateur passe de 3,6 à 12,8 millions "
    "de paramètres, d'où un entraînement environ 5 fois plus lent.")
header(s, "Énoncé · Q6", "Architecture et fonctions de perte")
# Dimensions reelles (verifiees en traversant les reseaux couche par couche).
def stage_row(y, title, stages, ops, final_fill):
    label(s, M, y - Inches(0.45), Inches(8), title, size=18, color=GOLD, align=PP_ALIGN.LEFT)
    bw, gap = Inches(1.12), Inches(0.33)
    for k, (res, ch) in enumerate(stages):
        x = M + k * (bw + gap)
        last = k == len(stages) - 1
        card(s, x, y, bw, Inches(0.95), fill=final_fill if last else STONE)
        text(s, x, y, bw, Inches(0.95), [res, [(ch, {"size": 11, "color": TEXT, "bold": False})]], size=15,
             bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, space_after=0)
        if not last:
            arrow(s, x + bw + Inches(0.03), y + Inches(0.33), Inches(0.27), Inches(0.28))
    text(s, M, y + Inches(1.02), Inches(8.2), Inches(0.35), ops, size=12, color=MUTED)


stage_row(Inches(2.1), "Générateur : agrandit le bruit",
          [("z", "100"), ("4×4", "512"), ("8×8", "256"), ("16×16", "128"), ("32×32", "64"), ("64×64", "3 · tanh")],
          "Convolutions transposées (pas 2) · BatchNorm · ReLU · sortie tanh dans [−1, 1]", RED)
stage_row(Inches(4.1), "Discriminateur : le chemin inverse",
          [("64×64", "3"), ("32×32", "64"), ("16×16", "128"), ("8×8", "256"), ("4×4", "512"), ("1", "logit")],
          "Convolutions (pas 2, sans pooling) · BatchNorm sauf 1re couche · LeakyReLU 0,2", RED)
chip(s, M, Inches(5.85), Inches(8.2), "128×128 : une couche de plus dans G et D · 12,8 M paramètres (G) contre 3,6 M",
     fill=CARD2, size=14, h=Inches(0.6))
card(s, Inches(9.1), Inches(1.8), Inches(3.63), Inches(4.65))
label(s, Inches(9.3), Inches(1.95), Inches(3.2), "Pertes", size=20, color=GOLD, align=PP_ALIGN.LEFT)
for j, (word, fill) in enumerate([("BCE", RED), ("D : vrai = 1, faux = 0", CARD2), ("G : −log D(G(z))", CARD2),
                                  ("Plan B : WGAN-GP", CARD2)]):
    chip(s, Inches(9.3), Inches(2.65) + j * Inches(0.9), Inches(3.23), word, fill=fill, size=15, h=Inches(0.65))

s = slide_(
    "Question 7 de l'énoncé — type de GAN et hyperparamètres.\n"
    "Nous avons choisi un DCGAN : c'est l'architecture de référence pour des images 64×64, simple à "
    "diagnostiquer et assez légère pour être entraînée sans GPU dédié. Des modèles comme StyleGAN donnent "
    "de meilleures images mais demandent des jours de calcul sur GPU.\n"
    "Hyperparamètres : ceux du papier DCGAN (Radford et al., 2015), qui sont un point de départ éprouvé — "
    "optimiseur Adam, taux d'apprentissage 0,0002, β1 = 0,5 (β1 plus bas que la valeur par défaut 0,9 "
    "pour stabiliser l'entraînement), lots de 128 images, vecteur de bruit de dimension 100, poids "
    "initialisés selon une loi normale d'écart-type 0,02. Augmentation : miroir horizontal aléatoire.\n"
    "Les valeurs affichées sont lues dans la configuration réellement utilisée pour l'entraînement.")
header(s, "Énoncé · Q7", "Un DCGAN, avec les réglages du papier")
params = [("Type", "DCGAN"), ("Optimiseur", "Adam"), ("Taux d'apprentissage", str(hp.get("lr", 0.0002)).replace(".", ",")),
          ("β1", str(hp.get("beta1", 0.5)).replace(".", ",")), ("Lot", str(hp.get("batch_size", 128))),
          ("Bruit z", str(hp.get("z_dim", 100))), ("Époques", str(hp.get("epochs", "—"))),
          ("Images d'entraînement", fmt(TRAIN["n_train"]) if "n_train" in TRAIN else "—")]
for j, (k_, v) in enumerate(params):
    r, c = divmod(j, 4)
    x, y = M + c * Inches(3.08), Inches(1.95) + r * Inches(2.3)
    card(s, x, y, Inches(2.9), Inches(2.0))
    big_number(s, x, y + Inches(0.3), Inches(2.9), v, k_, size=34 if len(v) < 8 else 26)

s = slide_(
    "Question 6 (suite) — l'entraînement.\n"
    "À gauche : les pertes du générateur et du discriminateur au fil des itérations. Elles oscillent au "
    "lieu de décroître, comme attendu pour un GAN : ce graphique sert surtout à repérer un effondrement "
    "(une perte de D qui tombe à zéro, par exemple).\n"
    "À droite : les images générées à partir des MÊMES vecteurs de bruit à plusieurs époques, ce qui "
    "permet de voir la progression réelle du générateur.\n"
    + ("" if TRAIN_DONE else "ÉTAT : entraînement en cours, les graphiques seront ajoutés automatiquement."))
header(s, "Énoncé · Q6", "Entraînement")
if TRAIN_DONE:
    chip(s, Inches(8.4), Inches(0.85), Inches(4.3), MODEL_LABEL, fill=RED if IS_V0 else CARD2, size=13, h=Inches(0.5))
result_figure(s, FIG_TR / "losses.png", M, Inches(1.75), Inches(5.9), Inches(4.8), "Pertes de G et D")
if TRAIN_DONE:
    chip(s, M, Inches(4.55), Inches(5.9), "Pertes qui oscillent : normal pour un GAN", fill=CARD2, size=15)
    chip(s, M, Inches(5.35), Inches(5.9), "D(x) au-dessus de D(G(z)) : D garde l'avantage", fill=RED, size=15)
result_figure(s, FIG_TR / "progression.png", Inches(6.75), Inches(1.75), Inches(5.95), Inches(4.8),
              "Mêmes z, époque après époque")

s = slide_(
    "Question 8 de l'énoncé — évaluation.\n"
    "Les métriques sont calculées entre des images générées et le jeu d'ÉVALUATION (10 % des œuvres, jamais "
    "vues pendant l'entraînement), en 64×64.\n"
    "• FID (Fréchet Inception Distance) : plus bas = mieux. Pour l'interpréter, on calcule aussi le FID "
    "entre deux moitiés de vraies peintures : c'est le meilleur score atteignable avec ce nombre d'images.\n"
    "• KID : même idée, plus fiable sur de petits échantillons.\n"
    "• Plus proches voisins : pour chaque image générée, la peinture d'entraînement la plus proche ; si elles "
    "sont identiques, le modèle recopie au lieu de créer.\n"
    + (f"\nRésultats ({MODEL_LABEL}, {TRAIN['epochs']} époques, {TRAIN['n_eval']} images d'évaluation) : "
       f"FID {TRAIN['fid']:.1f} contre {TRAIN['fid_real']:.1f} entre deux ensembles de vraies peintures de même taille ; "
       f"KID {TRAIN['kid']:.3f} ± {TRAIN['kid_std']:.3f} (référence réel/réel : {TRAIN['kid_real']:.3f}). "
       "L'écart entre le FID du générateur et le FID de référence mesure le chemin restant."
       if TRAIN_DONE else "ÉTAT : métriques en attente de la fin de l'entraînement."))
header(s, "Énoncé · Q8", "Évaluation du générateur")
if TRAIN_DONE:
    chip(s, Inches(8.4), Inches(0.85), Inches(4.3), MODEL_LABEL, fill=RED if IS_V0 else CARD2, size=13, h=Inches(0.5))
metrics = [("FID", f"{TRAIN['fid']:.1f}".replace(".", ",") if "fid" in TRAIN else "…", "généré vs réel", RED),
           ("FID réf.", f"{TRAIN['fid_real']:.1f}".replace(".", ",") if "fid_real" in TRAIN else "…", "réel vs réel", CARD2),
           ("KID", f"{TRAIN['kid']:.3f}".replace(".", ",") if "kid" in TRAIN else "…", "généré vs réel", CARD2)]
for j, (name, value, sub, fill) in enumerate(metrics):
    y = Inches(1.85) + j * Inches(1.55)
    card(s, M, y, Inches(4.2), Inches(1.35), fill=fill)
    text(s, M + Inches(0.3), y, Inches(1.8), Inches(1.35), [name, [(sub, {"size": 13, "color": MUTED if fill != RED else TEXT, "bold": False})]],
         size=20, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    text(s, M + Inches(2.0), y, Inches(2.0), Inches(1.35), value, size=36, bold=True, font=HEAD,
         color=GOLD if fill != RED else TEXT, align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)
result_figure(s, FIG_TR / "samples.png", Inches(5.2), Inches(1.75), Inches(3.7), Inches(4.8), "Images générées")
result_figure(s, FIG_TR / "nearest_neighbors.png", Inches(9.05), Inches(1.75), Inches(3.68), Inches(4.8),
              "Générée / plus proche réelle")

CMP_JSON = ROOT / "reports" / "training" / "comparison.json"
if TRAIN_DONE and not IS_V0 and CMP_JSON.exists():
    cmp = json.loads(CMP_JSON.read_text(encoding="utf-8"))
    ref, v0, v1 = cmp["reference (reel vs reel)"], cmp["dcgan_64_v0"], cmp["dcgan_64"]
    gain = 1 - v1["fid"] / v0["fid"]
    s = slide_(
        "Question 8 (suite) — la démarche v0 puis modèle final, et ce qu'elle prouve.\n"
        "La baseline v0 (30 époques sur 2 857 images) montrait un damier répété, un début de mode collapse et un "
        "discriminateur dominant. Notre hypothèse : il manquait des données et des itérations. Nous avons donc "
        "entraîné le modèle final avec 2,4 fois plus d'images et 7 fois plus d'itérations, SANS changer "
        "l'architecture, pour savoir exactement ce qui améliore le résultat.\n"
        f"Pour comparer honnêtement, les deux modèles sont évalués sur les MÊMES {cmp['n_eval']} peintures "
        f"d'évaluation : FID {v0['fid']:.0f} pour la v0, {v1['fid']:.0f} pour le modèle final, soit "
        f"{gain:.0%} de moins ; KID {v0['kid']:.2f} puis {v1['kid']:.2f}, divisé par deux. L'hypothèse est confirmée.\n"
        f"Le meilleur score atteignable avec ce nombre d'images est d'environ {ref['fid']:.0f} (FID entre deux "
        "ensembles de vraies peintures) : il reste donc de la marge. Les images ne montrent plus de damier ni de "
        "mode collapse ; on voit une matière picturale (ocres, clairs-obscurs, ciels) mais aucun objet "
        "identifiable, et un nouvel artefact : des traînées claires en diagonale. Les plus proches voisins sont "
        "des peintures différentes : le modèle crée, il ne recopie pas.")
    header(s, "Énoncé · Q8", "De la v0 au modèle final")
    chip(s, Inches(8.4), Inches(0.85), Inches(4.3), f"même évaluation · {cmp['n_eval']} peintures", fill=CARD2, size=13, h=Inches(0.5))
    data = CategoryChartData()
    data.categories = ["Baseline v0", "Modèle final", "Réel vs réel"]
    data.add_series("FID", [v0["fid"], v1["fid"], ref["fid"]])
    chart = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, M, Inches(1.75), Inches(6.6), Inches(4.9), data).chart
    dark_chart(chart, size=14)
    chart.has_legend = False
    plot = chart.plots[0]
    plot.gap_width = 70
    series = plot.series[0]
    for point, value, color in zip(series.points, [v0["fid"], v1["fid"], ref["fid"]], (STONE, RED, GOLD)):
        point.format.fill.solid()
        point.format.fill.fore_color.rgb = color
        tf = point.data_label.text_frame
        tf.text = f"FID {value:.0f}"
        run_ = tf.paragraphs[0].runs[0]
        run_.font.size, run_.font.bold, run_.font.color.rgb = Pt(16), True, TEXT
        point.data_label.position = XL_LABEL_POSITION.OUTSIDE_END
    chart.value_axis.visible = False
    chart.value_axis.has_major_gridlines = False
    chart.category_axis.tick_labels.font.color.rgb = TEXT
    chart.category_axis.format.line.color.rgb = STONE
    big_number(s, Inches(7.4), Inches(1.95), Inches(2.6), f"−{gain:.0%}".replace("\u00a0", ""), "FID", color=GOLD, size=44)
    big_number(s, Inches(10.1), Inches(1.95), Inches(2.6), "÷2", "KID", color=GOLD, size=44)
    for j, (word, fill) in enumerate([("Damier et collapse disparus", CARD2), ("Pas de recopie", CARD2),
                                      ("Traînées en diagonale", RED), ("Marge jusqu'à la référence", RED)]):
        chip(s, Inches(7.4), Inches(3.75) + j * Inches(0.75), Inches(5.3), word, fill=fill, size=16, h=Inches(0.6))

if TRAIN_DONE and IS_V0:
    hist = json.loads((ROOT / "runs" / "dcgan_64_v0" / "history.json").read_text(encoding="utf-8"))
    last_epoch = [h for h in hist if h["epoch"] == hist[-1]["epoch"]]
    d_real = np.mean([h["d_real"] for h in last_epoch])
    d_fake = np.mean([h["d_fake"] for h in last_epoch])
    s = slide_(
        "Question 8 (suite) — ce que nous observons sur la baseline v0, et ce que nous en concluons.\n"
        "• Un motif en damier se répète sur presque toutes les images : c'est un artefact connu des "
        "convolutions transposées en début d'entraînement.\n"
        "• Le même motif « fleuri » revient pour des bruits différents : c'est un début de mode collapse.\n"
        f"• Le discriminateur domine : il répond « vraie » à {d_real:.2f} pour les vraies images et à "
        f"{d_fake:.2f} pour les générées, loin de l'équilibre à 0,5.\n"
        f"• Seulement {TRAIN['iterations']} itérations (30 époques de {TRAIN['n_train']} images) : un DCGAN en "
        "demande généralement plusieurs milliers.\n"
        "• Point positif : la structure ciel clair / sol sombre et les tons bruns des tableaux anciens "
        "apparaissent, et les plus proches voisins sont des peintures différentes : pas de recopie.\n"
        "Hypothèse pour la v1 : plus de données (le jeu complet, environ 7 200 images) et plus d'itérations "
        "(environ 100 époques), sans changer l'architecture, pour savoir ce qui améliore le résultat. Si le "
        "damier persiste, nous remplacerons les convolutions transposées par un agrandissement suivi d'une "
        "convolution ; si D domine toujours, nous adoucirons ses étiquettes (0,9 au lieu de 1).")
    header(s, "Énoncé · Q8", "Diagnostic de la baseline v0")
    chip(s, Inches(8.4), Inches(0.85), Inches(4.3), MODEL_LABEL, fill=RED, size=13, h=Inches(0.5))
    label(s, M, Inches(1.8), Inches(6), "Observations", size=22, color=GOLD, align=PP_ALIGN.LEFT)
    for j, (word, fill) in enumerate([("Damier répété", RED), ("Début de mode collapse", RED),
                                      (f"D domine : {d_real:.2f} / {d_fake:.2f}".replace(".", ","), CARD2),
                                      (f"{TRAIN['iterations']} itérations seulement", CARD2),
                                      ("Ciel / sol appris · pas de recopie", CARD2)]):
        chip(s, M, Inches(2.55) + j * Inches(0.82), Inches(5.6), word, fill=fill, size=17, h=Inches(0.65))
    card(s, Inches(6.9), Inches(1.8), Inches(5.83), Inches(4.75))
    label(s, Inches(7.2), Inches(1.95), Inches(5.3), "Hypothèse v1", size=22, color=GOLD, align=PP_ALIGN.LEFT)
    big_number(s, Inches(6.9), Inches(2.7), Inches(2.9), "×2,5", "images", size=44)
    big_number(s, Inches(9.8), Inches(2.7), Inches(2.9), "×8", "itérations", size=44)
    chip(s, Inches(7.2), Inches(4.75), Inches(5.2), "même architecture : une variable à la fois", fill=CARD2, size=15)
    chip(s, Inches(7.2), Inches(5.55), Inches(5.2), "plan B : upsample + conv · étiquettes 0,9", fill=CARD2, size=15)

# ================================================================ Jeu homogene : portraits (cours)
P64_JSON = ROOT / "reports" / "portraits_64" / "metrics.json"
P128_JSON = ROOT / "reports" / "portraits_128" / "metrics.json"
DL16 = ROOT / "reports" / "portraits" / "figures" / "dataloader_16.png"
if P64_JSON.exists() and DL16.exists():
    p64 = json.loads(P64_JSON.read_text(encoding="utf-8"))
    p128 = json.loads(P128_JSON.read_text(encoding="utf-8")) if P128_JSON.exists() else None
    v1m = TRAIN

    s = slide_(
        "Notre cours de traitement d'images pose une règle d'or : un dataset homogène (un style, un format, un "
        "cadrage), sinon le GAN produit des images floues ou fait un mode collapse. Notre premier jeu mélangeait "
        "10 genres et 7 siècles, et nos images restaient abstraites : nous avons donc testé la règle.\n"
        "Nous avons constitué un jeu de portraits : 4 068 sélectionnés (le maximum disponible dans le catalogue "
        "après nos filtres), 3 737 téléchargés, dont 3 347 pour l'entraînement et 390 pour l'évaluation, toujours "
        "séparés par œuvre.\n"
        "Nous avons appliqué le pipeline du cours à la lettre : dossier dédié en 128×128, classe Dataset qui lit "
        "les images à la demande, transformations Resize, CenterCrop, RandomHorizontalFlip, ToTensor et Normalize "
        "dans cet ordre, puis DataLoader.\n"
        "Avant d'entraîner, nous avons vérifié 16 images à la sortie du DataLoader, comme le recommande le cours : "
        "format (16, 3, 128, 128), valeurs entre −1 et 1, cadrage centré sans déformation. Limite : le jeu reste "
        "imparfaitement homogène (portraits en pied, de groupe, équestres).")
    header(s, "Cours · traitement d'images", "Un jeu homogène : les portraits")
    figure(s, DL16, M, Inches(1.75), Inches(12.1), Inches(3.4), "16 images à la sortie du DataLoader (après pipeline)")
    for j, (word, fill) in enumerate([("3 737 portraits", RED), ("Dossier 128×128", CARD2),
                                      ("Resize → Crop → Flip → Tensor → Normalize", CARD2), ("[−1, 1] vérifié", CARD2)]):
        chip(s, M + j * Inches(3.08), Inches(5.75), Inches(2.95), word, fill=fill, size=13, h=Inches(0.6))

    gap_v1 = v1m["fid"] / v1m["fid_real"]
    gap_p = p64["fid"] / p64["fid_real"]
    s = slide_(
        "Résultat du modèle portraits en 64×64 : même architecture, mêmes hyperparamètres et autant d'itérations "
        "(environ 4 700) que le modèle tous genres. Seul le jeu de données change.\n"
        f"FID {p64['fid']:.0f} contre {p64['fid_real']:.0f} entre vraies peintures, soit {gap_p:.1f} fois la référence ; "
        f"le modèle tous genres était à {v1m['fid']:.0f} contre {v1m['fid_real']:.0f}, soit {gap_v1:.1f} fois. "
        f"KID {p64['kid']:.3f} contre {v1m['kid']:.3f}. Les deux modèles ne sont pas évalués sur les mêmes peintures : "
        "on compare donc l'écart à la référence de chacun.\n"
        "La différence est surtout visible : presque chaque image est un portrait reconnaissable (silhouette "
        "centrée, visage clair, buste, fond sombre, perruques ou cols blancs). Les visages restent flous, ce qui est "
        "normal en 64×64. Les plus proches voisins sont des tableaux différents : pas de recopie.\n"
        "Conclusion : la règle d'or du cours est vérifiée sur nos données.")
    header(s, "Cours · traitement d'images", "Portraits 64×64 : la règle se vérifie")
    figure(s, ROOT / "reports" / "portraits_64" / "figures" / "samples.png", M, Inches(1.75), Inches(4.9), Inches(4.9),
           "Portraits générés (64×64)")
    rows = [("Tous genres (v1)", v1m["fid"], v1m["fid_real"], CARD2), ("Portraits 64×64", p64["fid"], p64["fid_real"], RED)]
    label(s, Inches(5.85), Inches(1.85), Inches(6.9), "FID   ·   référence   ·   écart", size=15, color=GOLD,
          align=PP_ALIGN.LEFT)
    for j, (name, f, ref, fill) in enumerate(rows):
        y = Inches(2.45) + j * Inches(1.05)
        card(s, Inches(5.85), y, Inches(6.88), Inches(0.9), fill=fill)
        text(s, Inches(6.1), y, Inches(3.0), Inches(0.9), name, size=17, bold=True, anchor=MSO_ANCHOR.MIDDLE)
        text(s, Inches(9.1), y, Inches(3.4), Inches(0.9), f"{f:.0f}  ·  {ref:.0f}  ·  ×{f / ref:.1f}".replace(".", ","),
             size=20, bold=True, font=HEAD, anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.RIGHT)
    for j, (word, fill) in enumerate([("Portraits reconnaissables", CARD2), ("Pas de recopie", CARD2),
                                      ("Visages encore flous", RED)]):
        chip(s, Inches(5.85), Inches(4.75) + j * Inches(0.65), Inches(6.88), word, fill=fill, size=15, h=Inches(0.55))

    s = slide_(
        "Deuxième règle du cours : commencer en 64×64, puis passer en 128×128 pour gagner en détails. Nous avons "
        "étendu le DCGAN d'une couche dans le générateur et dans le discriminateur (12,8 millions de paramètres "
        "pour le générateur, contre 3,6 en 64×64).\n"
        "Sur notre processeur, une époque en 128×128 prend 5 min 20 contre 1 min en 64×64. Pour tenir les délais, "
        "nous l'avons limité à 100 époques (environ 2 600 itérations, contre 4 700 en 64×64) : la comparaison est "
        "donc moins équitable, ce qui est une limite à garder en tête.\n"
        + (f"Résultat : FID {p128['fid']:.0f} contre {p128['fid_real']:.0f} entre vraies peintures, soit "
           f"{p128['fid'] / p128['fid_real']:.1f} fois la référence. Les FID en 64 et en 128 ne se comparent pas "
           "directement (les images sont redimensionnées différemment avant Inception) : on compare l'écart et les images.\n"
           f"Conclusion : avec {p128['iterations']} itérations seulement, le 128×128 reste moins bon que le 64×64 "
           f"(×{p128['fid'] / p128['fid_real']:.1f} contre ×{p64['fid'] / p64['fid_real']:.1f}) : traînées horizontales, couleurs "
           "saturées, silhouettes moins lisibles. Un modèle 3,5 fois plus gros demande plus d'entraînement, pas moins. "
           "Nous gardons le 64×64 comme modèle principal ; le 128×128 est la prochaine étape, sur GPU."
           if p128 else "ÉTAT : entraînement en cours ; résultats ajoutés automatiquement à la fin."))
    header(s, "Cours · traitement d'images", "Passer en 128×128")
    result_figure(s, ROOT / "reports" / "portraits_128" / "figures" / "samples.png", M, Inches(1.75), Inches(4.9),
                  Inches(4.9), "Portraits générés (128×128)")
    big_number(s, Inches(5.85), Inches(1.95), Inches(3.4), "×4", "pixels par image", size=44)
    big_number(s, Inches(9.3), Inches(1.95), Inches(3.4), "5 min 20", "par époque (CPU)", color=RED, size=40)
    chips128 = [("100 époques au lieu de 180", RED), ("Une couche de plus dans G et D", CARD2)]
    if p128:
        chips128.insert(0, (f"FID {p128['fid']:.0f} · référence {p128['fid_real']:.0f} · ×{p128['fid'] / p128['fid_real']:.1f}"
                            .replace(".", ","), CARD2))
        chips128.append((f"Moins bon que le 64×64 (×{p64['fid'] / p64['fid_real']:.1f})".replace(".", ","), RED))
    for j, (word, fill) in enumerate(chips128):
        chip(s, Inches(5.85), Inches(3.7) + j * Inches(0.7), Inches(6.88), word, fill=fill, size=15, h=Inches(0.58))

# ================================================================ Equipe
s = slide_(
    "À compléter par l'équipe : les outils utilisés (dépôt GitHub, tableau Trello ou Azure DevOps), "
    "les rituels (fréquence des points d'équipe, revues de code) et la répartition des tâches entre "
    "les membres.")
header(s, "Équipe", "Organisation")
for k, (title, words) in enumerate([("Outils", ["GitHub", "[Trello / Azure]"]),
                                    ("Rituels", ["[Points d'équipe]", "[Revues]"]),
                                    ("Rôles", ["[Membre 1]", "[Membre 2]", "[Membre 3]"])]):
    x = M + k * (cw + Inches(0.3))
    card(s, x, Inches(1.85), cw, Inches(4.9))
    circle(s, x + Inches(0.35), Inches(2.15), Inches(0.75), str(k + 1), size=18)
    text(s, x + Inches(1.3), Inches(2.25), cw - Inches(1.5), Inches(0.6), title, size=22, bold=True,
         font=HEAD, color=GOLD)
    for j, word in enumerate(words):
        chip(s, x + Inches(0.35), Inches(3.3) + j * Inches(0.8), cw - Inches(0.7), word, size=17)

s = slide_(
    "Question 9 de l'énoncé — conclusion.\n"
    "Acquis : le catalogue est analysé, ses biais sont identifiés et corrigés par la sélection ; le jeu "
    "d'images est reproductible, contrôlé et testé ; le DCGAN est entraîné et évalué avec un protocole "
    "sans fuite.\n"
    "Difficultés : un serveur lent et irrégulier (blocage puis correctif), des données hétérogènes, un fort "
    "déséquilibre des genres, et un faux signal de doublons repéré grâce aux visualisations.\n"
    "Pistes : WGAN-GP si le mode collapse apparaît, passage en 128×128, GAN conditionnel pour choisir le "
    "genre, augmentation adaptative (ADA), et l'application Streamlit pour rendre le générateur accessible.")
header(s, "Énoncé · Q9", "Conclusion")
cols = [("Acquis", ["Catalogue analysé", "Pipeline du cours", "Portraits reconnaissables"], CARD2),
        ("Difficultés", ["Serveur lent", "Faux doublons", "Genres déséquilibrés"], RED),
        ("Pistes", ["WGAN-GP", "GAN conditionnel", "GPU pour le 128×128"], CARD2)]
for k, (title, words, fill) in enumerate(cols):
    x = M + k * (cw + Inches(0.3))
    text(s, x, Inches(1.95), cw, Inches(0.6), title, size=26, bold=True, font=HEAD, color=GOLD)
    for j, word in enumerate(words):
        chip(s, x, Inches(2.8) + j * Inches(0.8), cw, word, fill=fill, size=17)
side = Inches(1.25)
for k, path in enumerate(safe(9, 6)):
    picture(s, path, M + k * (side + Inches(0.11)), Inches(5.6), side)

# ================================================================ Application Streamlit
if (ROOT / "app.py").exists() and (ROOT / "reports" / "app" / "preview_seed42.png").exists():
    s = slide_(
        "GANisme à l'infini : l'application Streamlit (fichier app.py) rend le générateur accessible à tous.\n"
        "On choisit le nombre de peintures (de 1 à 36) et une graine : la même graine redonne toujours les "
        "mêmes peintures, ce qui permet de retrouver une création. Le bouton « Nouvelles peintures » tire une "
        "graine au hasard. Les images, générées en 64×64 pixels, sont agrandies sans flou, et la grille peut "
        "être téléchargée en PNG.\n"
        "L'application charge un fichier léger (models/generator.pt, environ 14 Mo) qui ne contient que les poids "
        "du générateur : elle fonctionne sans les checkpoints d'entraînement. Le modèle est chargé une seule "
        "fois grâce au cache de Streamlit. L'interface est couverte par des tests automatiques (AppTest).\n"
        "Lancement : streamlit run app.py")
    header(s, "GANisme à l'infini", "L'application Streamlit")
    figure(s, ROOT / "reports" / "app" / "preview_seed42.png", M, Inches(1.75), Inches(5.0), Inches(4.9),
           "Grille réelle produite par l'application (graine 42)")
    for j, word in enumerate(["Graine reproductible", "1 à 36 peintures", "Agrandissement sans flou",
                              "Téléchargement PNG"]):
        chip(s, Inches(6.0), Inches(1.9) + j * Inches(0.8), Inches(6.7), word, fill=CARD2, size=17, h=Inches(0.62))
    terminal(s, Inches(6.0), Inches(5.2), Inches(6.7), Inches(1.3), ["streamlit run app.py"], "Lancement")

# ================================================================ Couverture de l'enonce
s = slide_(
    "Tableau de couverture de l'énoncé : chaque question est traitée dans la présentation (diapositive "
    "indiquée) et dans les notebooks. Les questions marquées « en cours » dépendent de l'entraînement du "
    "modèle et seront complétées dès qu'il sera terminé.")
header(s, "Énoncé", "Couverture de l'énoncé")
done_q = "✓"
todo_q = "✓" if TRAIN_DONE else "en cours"
items_left = [(f"Veille Q{i}", done_q) for i in range(1, 16)]
items_right = [("1 · Dataset", done_q), ("2 · Peintures", done_q), ("3 · Analyse", done_q),
               ("4 · Scraping", "✓" if n_pend == 0 else "en cours"), ("5 · Visualisation", done_q),
               ("6 · Architecture, pertes", done_q), ("6 · Entraînement", todo_q), ("7 · Type, hyperparamètres", done_q),
               ("8 · Évaluation", todo_q), ("9 · Conclusion", todo_q), ("Streamlit", "✓" if (ROOT / "app.py").exists() else "à faire")]
for col, items in enumerate((items_left, items_right)):
    per_row = 3 if col == 0 else 2
    w_ = Inches(1.95) if col == 0 else Inches(2.85)
    x0 = M if col == 0 else Inches(6.85)
    for j, (name, state) in enumerate(items):
        r, c = divmod(j, per_row)
        x, y = x0 + c * (w_ + Inches(0.08)), Inches(1.85) + r * Inches(0.82)
        fill = CARD2 if state == "✓" else RED
        chip(s, x, y, w_, f"{name}  {state}", fill=fill, size=13 if col == 0 else 14, h=Inches(0.66))


prs.core_properties.title = "GANisme"
OUT.parent.mkdir(parents=True, exist_ok=True)
prs.save(OUT)
print(f"{OUT} : {len(prs.slides)} diapositives ; etat {stamp}")
