"""Genere reports/presentation/GANisme.pptx a partir des resultats reels du projet.

Les chiffres viennent de reports/eda/summary.json et de data/processed/manifest.csv,
les images de data/processed/images_256/. Relancer le script met la presentation a
jour (par exemple apres la fin du telechargement ou l'entrainement du GAN) :

    .venv/Scripts/python.exe reports/presentation/build_presentation.py
"""

import json
import random
from datetime import datetime
from pathlib import Path

import pandas as pd
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "reports" / "presentation" / "GANisme.pptx"
EDA = json.loads((ROOT / "reports" / "eda" / "summary.json").read_text(encoding="utf-8"))
FIG = ROOT / "reports" / "eda" / "figures"
MANIFEST = pd.read_csv(ROOT / "data" / "processed" / "manifest.csv")
IMAGES = ROOT / "data" / "processed" / "images_256"

# Palette "musee" : encre chaude, rouge sang-de-boeuf, or (sur fond sombre uniquement).
INK = RGBColor(0x1E, 0x1B, 0x18)
OXBLOOD = RGBColor(0x8C, 0x2F, 0x39)
GOLD = RGBColor(0xD4, 0xA8, 0x4B)
GREY = RGBColor(0x5F, 0x5A, 0x55)
STONE = RGBColor(0xA8, 0xA3, 0x9D)
CARD = RGBColor(0xF2, 0xF1, 0xEF)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
HEAD, BODY = "Cambria", "Calibri"

W, H = Inches(13.333), Inches(7.5)
MARGIN = Inches(0.6)

prs = Presentation()
prs.slide_width, prs.slide_height = W, H
BLANK = prs.slide_layouts[6]
GENRE_FR = {
    "religious": "religieux", "portrait": "portrait", "landscape": "paysage",
    "mythological": "mythologique", "genre": "scène de genre", "still-life": "nature morte",
    "other": "autre", "historical": "historique", "interior": "intérieur", "study": "étude",
}


def fmt(n):
    return f"{n:,}".replace(",", " ")


def pct(x, digits=0):
    return f"{x * 100:.{digits}f}".replace(".", ",") + " %"


# ---------------------------------------------------------------- primitives

def new_slide(dark=False, notes=None):
    slide = prs.slides.add_slide(BLANK)
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = INK if dark else WHITE
    if notes:
        slide.notes_slide.notes_text_frame.text = notes
    return slide


def _bullet(paragraph, level=0):
    ppr = paragraph._p.get_or_add_pPr()
    indent = Emu(Inches(0.22))
    ppr.set("marL", str(int(indent * (level + 1))))
    ppr.set("indent", str(-int(indent)))
    for tag in ("a:buNone", "a:buChar", "a:buAutoNum"):
        for old in ppr.findall(qn(tag)):
            ppr.remove(old)
    bu = ppr.makeelement(qn("a:buChar"), {"char": "•"})
    ppr.append(bu)


def text(slide, x, y, w, h, content, size=16, color=INK, bold=False, font=BODY,
         align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, bullets=False, space_after=6, name=None):
    """Zone de texte. ``content`` : chaine, ou liste de paragraphes ; un paragraphe est
    une chaine ou une liste de runs ``(texte, {bold, color, size, italic, font})``."""
    box = slide.shapes.add_textbox(x, y, w, h)
    if name:
        box.name = name
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    paragraphs = content if isinstance(content, list) else [content]
    for i, para in enumerate(paragraphs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space_after)
        if bullets:
            _bullet(p)
        runs = para if isinstance(para, list) else [(para, {})]
        for run_text, opts in runs:
            run = p.add_run()
            run.text = run_text
            f = run.font
            f.name = opts.get("font", font)
            f.size = Pt(opts.get("size", size))
            f.bold = opts.get("bold", bold)
            f.italic = opts.get("italic", False)
            f.color.rgb = opts.get("color", color)
    return box


def box(slide, x, y, w, h, fill=CARD, radius=0.06, line=None, name=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    shape.adjustments[0] = radius
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
        shape.line.width = Pt(1.25)
    shape.shadow.inherit = False
    if name:
        shape.name = name
    return shape


def circle(slide, x, y, d, label, fill=OXBLOOD, color=WHITE, size=16):
    shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, x, y, d, d)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.fill.background()
    shape.shadow.inherit = False
    tf = shape.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    run = p.add_run()
    run.text = label
    run.font.name, run.font.size, run.font.bold = BODY, Pt(size), True
    run.font.color.rgb = color
    return shape


def arrow(slide, x, y, w, h, fill=STONE):
    shape = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, x, y, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.fill.background()
    shape.shadow.inherit = False
    return shape


def header(slide, kicker, title):
    """Surtitre (section) + titre, a la meme position sur toutes les diapositives de contenu."""
    text(slide, MARGIN, Inches(0.45), Inches(12), Inches(0.3), kicker.upper(), size=12,
         color=OXBLOOD, bold=True, name="Surtitre")
    text(slide, MARGIN, Inches(0.75), Inches(12.1), Inches(0.8), title, size=32, bold=True,
         font=HEAD, name="Titre")


def footer(slide, label, dark=False):
    text(slide, MARGIN, Inches(7.05), Inches(8), Inches(0.25), label, size=10,
         color=STONE if dark else GREY, name="Pied de page")


def table(slide, x, y, w, rows, col_widths, size=13, header_fill=INK, row_h=0.42):
    shape = slide.shapes.add_table(len(rows), len(rows[0]), x, y, w, Inches(row_h * len(rows)))
    tbl = shape.table
    for j, cw in enumerate(col_widths):
        tbl.columns[j].width = Inches(cw)
    for i, row in enumerate(rows):
        for j, value in enumerate(row):
            cell = tbl.cell(i, j)
            cell.margin_left = cell.margin_right = Inches(0.1)
            cell.margin_top = cell.margin_bottom = Inches(0.05)
            cell.fill.solid()
            cell.fill.fore_color.rgb = header_fill if i == 0 else (CARD if i % 2 else WHITE)
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            run = p.add_run()
            run.text = value
            run.font.name, run.font.size = BODY, Pt(size)
            run.font.bold = i == 0 or j == 0
            run.font.color.rgb = WHITE if i == 0 else INK
    return tbl


def style_chart(chart, size=12):
    chart.font.name = BODY
    chart.font.size = Pt(size)
    chart.font.color.rgb = GREY


def processed_images(n, seed=0, genre=None):
    ok = MANIFEST[MANIFEST["status"] == "processed"]
    if genre:
        ok = ok[ok["TYPE"] == genre]
    paths = [IMAGES / Path(p).name for p in ok["local_path"]]
    paths = [p for p in paths if p.exists()]
    random.Random(seed).shuffle(paths)
    return paths[:n]


def divider(number, title, subtitle, notes):
    slide = new_slide(dark=True, notes=notes)
    text(slide, MARGIN, Inches(2.4), Inches(3), Inches(1.4), number, size=96, bold=True,
         font=HEAD, color=GOLD)
    text(slide, MARGIN, Inches(3.85), Inches(11), Inches(0.9), title, size=40, bold=True,
         font=HEAD, color=WHITE)
    text(slide, MARGIN, Inches(4.75), Inches(10), Inches(0.6), subtitle, size=18, color=STONE)
    strip = processed_images(6, seed=int(number))
    side = Inches(1.15)
    for k, path in enumerate(strip):
        slide.shapes.add_picture(str(path), W - MARGIN - (6 - k) * (side + Inches(0.12)) + Inches(0.12),
                                 Inches(0.6), side, side)
    return slide


# ---------------------------------------------------------------- donnees
status = MANIFEST["status"].value_counts()
n_processed = int(status.get("processed", 0))
n_failed = int(status.get("failed", 0))
n_skipped = int(status.get("skipped", 0))
n_pending = int(status.get("pending", 0))
n_selected = len(MANIFEST)
download_done = n_pending == 0
stamp = datetime.now().strftime("%d/%m/%Y %H:%M")
types = pd.Series(EDA["types"])
n_paint = EDA["n_paintings"]

# ================================================================ 1. Titre
slide = new_slide(dark=True, notes=(
    "Présentation du projet GANisme : générer des peintures originales avec un réseau "
    "antagoniste génératif, entraîné sur le catalogue de la Web Gallery of Art."))
text(slide, MARGIN, Inches(2.2), Inches(6.2), Inches(1.4), "GANisme", size=72, bold=True,
     font=HEAD, color=WHITE)
text(slide, MARGIN, Inches(3.6), Inches(5.9), Inches(1.2),
     "Générer des peintures avec un réseau antagoniste génératif", size=24, color=GOLD, font=HEAD)
text(slide, MARGIN, Inches(5.2), Inches(5.9), Inches(0.9),
     ["La Plateforme · Projet Deep Learning", "Équipe : [à compléter]"], size=14, color=STONE)
grid = processed_images(20, seed=7)
side, gap = Inches(1.2), Inches(0.08)
for k, path in enumerate(grid):
    r, c = divmod(k, 4)
    slide.shapes.add_picture(str(path), Inches(7.55) + c * (side + gap), Inches(0.55) + r * (side + gap),
                             side, side)
text(slide, Inches(7.55), Inches(7.0), Inches(5.5), Inches(0.25),
     "Images réelles du jeu d'entraînement (Web Gallery of Art)", size=10, color=STONE)

# ================================================================ 2. Plan
slide = new_slide(notes="Le plan suit les attendus de l'énoncé.")
header(slide, "Plan", "Du catalogue au générateur de peintures")
parts = [
    ("1", "Veille scientifique", "Les 15 questions sur les GAN"),
    ("2", "Le catalogue", "Analyse des 52 867 œuvres"),
    ("3", "Le jeu d'images", "Sélection, scraping, prétraitement"),
    ("4", "Le modèle", "Architecture DCGAN, résultats"),
    ("5", "Équipe et conclusion", "Organisation, difficultés, suite"),
]
for k, (num, title, sub) in enumerate(parts):
    y = Inches(1.85) + k * Inches(1.0)
    circle(slide, MARGIN, y, Inches(0.7), num, size=20)
    text(slide, MARGIN + Inches(0.95), y + Inches(0.02), Inches(5), Inches(0.4), title, size=20,
         bold=True, font=HEAD)
    text(slide, MARGIN + Inches(0.95), y + Inches(0.4), Inches(5), Inches(0.3), sub, size=14, color=GREY)
pics = processed_images(4, seed=3)
for k, path in enumerate(pics):
    r, c = divmod(k, 2)
    slide.shapes.add_picture(str(path), Inches(7.6) + c * Inches(2.55), Inches(1.85) + r * Inches(2.55),
                             Inches(2.4), Inches(2.4))

# ================================================================ 3-10. Veille
divider("01", "Veille scientifique", "Comprendre les GAN avant de les entraîner",
        "Première partie : les réponses aux 15 questions de veille de l'énoncé.")

# Q1-Q2 : probleme + fonctionnement (schema)
slide = new_slide(notes=(
    "Q1 : le GAN répond à l'apprentissage non supervisé d'un modèle génératif. "
    "Q2 : deux réseaux s'affrontent ; le générateur apprend à tromper le discriminateur."))
header(slide, "Veille · Q1–Q2", "Deux réseaux qui s'affrontent pour apprendre à créer")
text(slide, MARGIN, Inches(1.75), Inches(4.6), Inches(4.8), [
    [("Q1 · Problème initial. ", {"bold": True}),
     ("Apprentissage non supervisé d'un modèle génératif : produire de nouvelles données qui "
      "suivent la distribution d'un jeu d'exemples, sans étiquettes.", {})],
    [("Ailleurs ? ", {"bold": True}),
     ("Efficace en image-à-image, super-résolution, augmentation de données, semi-supervisé. "
      "Moins adapté au texte : l'échantillonnage discret bloque le gradient.", {})],
    [("Q2 · Principe. ", {"bold": True}),
     ("Jeu à somme nulle : le générateur fabrique des faux, le discriminateur les démasque. "
      "Chacun s'améliore contre l'autre jusqu'à ce que les faux soient indiscernables.", {})],
], size=15, space_after=12)
# schema : z -> G -> image generee ; vraie peinture et image generee -> D -> vrai / faux
top, low = Inches(1.85), Inches(3.75)          # lignes "vraie peinture" et "image generee"
img = Inches(1.1)
x_z, x_g, x_img, x_d = Inches(5.75), Inches(7.6), Inches(9.45), Inches(11.0)
bw, bh = Inches(1.4), Inches(0.95)
mid = low + (img - bh) / 2
box(slide, x_z, mid, bw, bh)
text(slide, x_z, mid, bw, bh, ["Bruit z", [("100 nombres aléatoires", {"size": 11, "color": GREY})]],
     size=15, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, space_after=0)
arrow(slide, x_z + bw + Inches(0.08), low + Inches(0.4), Inches(0.3), Inches(0.3))
box(slide, x_g, mid, bw, bh, fill=OXBLOOD)
text(slide, x_g, mid, bw, bh, "Générateur G", size=15, bold=True, color=WHITE,
     align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
arrow(slide, x_g + bw + Inches(0.08), low + Inches(0.4), Inches(0.3), Inches(0.3))
real, fake = processed_images(1, seed=12), processed_images(1, seed=11)
slide.shapes.add_picture(str(real[0]), x_img, top, img, img)
text(slide, x_img - Inches(0.2), top + img + Inches(0.05), img + Inches(0.4), Inches(0.3), "vraie peinture",
     size=11, color=GREY, align=PP_ALIGN.CENTER)
slide.shapes.add_picture(str(fake[0]), x_img, low, img, img)
text(slide, x_img - Inches(0.2), low + img + Inches(0.05), img + Inches(0.4), Inches(0.3), "image générée",
     size=11, color=GREY, align=PP_ALIGN.CENTER)
for row in (top, low):
    arrow(slide, x_img + img + Inches(0.08), row + Inches(0.4), Inches(0.3), Inches(0.3))
d_top = top
d_h = low + img - top
box(slide, x_d + Inches(0.05), d_top, Inches(1.65), d_h, fill=INK)
text(slide, x_d + Inches(0.05), d_top, Inches(1.65), d_h,
     ["Discriminateur D", [("« vrai ou faux ? »", {"size": 13, "color": GOLD, "bold": False})]],
     size=15, bold=True, color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
text(slide, x_z, Inches(5.25), Inches(6.9), Inches(0.4),
     "Le gradient de D indique à G comment mieux le tromper.", size=14, color=GREY)
text(slide, x_z, Inches(5.8), Inches(6.9), Inches(0.5),
     [[("min", {"bold": True}), ("G", {"size": 10}), (" max", {"bold": True}), ("D", {"size": 10}),
       ("  E[log D(x)] + E[log(1 − D(G(z)))]", {})]], size=17, font=HEAD, color=OXBLOOD)
footer(slide, "Goodfellow et al., Generative Adversarial Nets, 2014")

# Q3-Q5
slide = new_slide(notes=(
    "Q3 : rôles et différence génératif / discriminant. Q4 : entrées. Q5 : densité implicite."))
header(slide, "Veille · Q3–Q5", "Générer n'est pas classer")
cards = [
    ("Q3 · Rôles", [
        [("Générateur : ", {"bold": True}), ("transforme du bruit en image réaliste.", {})],
        [("Discriminateur : ", {"bold": True}), ("classifieur binaire vrai / faux, qui sert de critique à G.", {})],
        [("Modèle discriminant : ", {"bold": True}), ("apprend p(y | x), une frontière entre classes.", {})],
        [("Modèle génératif : ", {"bold": True}), ("apprend p(x) ; il peut produire de nouveaux x.", {})],
    ]),
    ("Q4 · Entrées", [
        [("G reçoit ", {"bold": True}), ("un vecteur de bruit z (ex. 100 valeurs tirées d'une loi normale), plus une étiquette pour un GAN conditionnel.", {})],
        [("D reçoit ", {"bold": True}), ("une image, vraie (jeu de données) ou générée, et sort la probabilité qu'elle soit vraie.", {})],
    ]),
    ("Q5 · Densité implicite", [
        [("Le GAN ne donne jamais la formule de p(x) ni la vraisemblance d'une image.", {})],
        [("La distribution n'existe qu'à travers l'échantillonnage z → G(z).", {})],
        [("À l'inverse, VAE et flows sont des modèles à densité explicite.", {"color": GREY})],
    ]),
]
cw = Inches(3.85)
for k, (title, body) in enumerate(cards):
    x = MARGIN + k * (cw + Inches(0.3))
    box(slide, x, Inches(1.8), cw, Inches(4.9))
    text(slide, x + Inches(0.3), Inches(2.05), cw - Inches(0.6), Inches(0.4), title, size=20,
         bold=True, font=HEAD, color=OXBLOOD)
    text(slide, x + Inches(0.3), Inches(2.65), cw - Inches(0.6), Inches(3.9), body, size=16,
         bullets=True, space_after=12)

# Q6 + Q9
slide = new_slide(notes=(
    "Q6 : rétropropagation alternée. Q9 : la convergence est un équilibre, pas un minimum ; "
    "les pertes oscillent, d'où le suivi visuel."))
header(slide, "Veille · Q6 et Q9", "Un entraînement alterné qui cherche un équilibre")
steps = [
    ("1", "Pas du discriminateur",
     "Un lot de vraies images (étiquette 1) et de fausses (étiquette 0). Perte d'entropie croisée "
     "binaire, rétropropagée dans D seul : G est figé."),
    ("2", "Pas du générateur",
     "G génère, D juge. Perte −log D(G(z)) : le gradient traverse D (figé) jusqu'à G, "
     "dont seuls les poids sont mis à jour."),
]
for k, (num, title, body) in enumerate(steps):
    y = Inches(1.85) + k * Inches(1.55)
    circle(slide, MARGIN, y, Inches(0.6), num)
    text(slide, MARGIN + Inches(0.85), y, Inches(5.1), Inches(0.4), title, size=18, bold=True, font=HEAD)
    text(slide, MARGIN + Inches(0.85), y + Inches(0.42), Inches(5.1), Inches(1.0), body, size=14)
text(slide, MARGIN, Inches(5.2), Inches(5.9), Inches(1.4),
     [[("Q6 · ", {"bold": True, "color": OXBLOOD}),
       ("On répète ces deux pas à chaque lot. La rétropropagation est classique, seule change la cible "
        "de la perte selon le réseau entraîné.", {})]], size=14)
box(slide, Inches(7.0), Inches(1.8), Inches(5.7), Inches(4.9))
text(slide, Inches(7.3), Inches(2.05), Inches(5.1), Inches(0.4), "Q9 · Convergence", size=20,
     bold=True, font=HEAD, color=OXBLOOD)
text(slide, Inches(7.3), Inches(2.65), Inches(5.1), Inches(3.9), [
    [("Équilibre théorique : ", {"bold": True}), ("les faux suivent la distribution réelle et D répond 0,5 partout.", {})],
    [("En pratique, ", {"bold": True}), ("les pertes oscillent au lieu de décroître : elles ne disent pas si G progresse.", {})],
    [("D trop fort : ", {"bold": True}), ("D(G(z)) ≈ 0, les gradients de G s'évanouissent.", {})],
    [("D trop faible : ", {"bold": True}), ("G ne reçoit plus de signal utile.", {})],
    [("Mode collapse : ", {"bold": True}), ("G ne produit plus que quelques variantes qui trompent D.", {})],
], size=14, bullets=True, space_after=9)

# Q7-Q8
slide = new_slide(notes=(
    "Q7 : aucune métrique unique ; on croise FID/KID, précision/rappel et inspection visuelle. "
    "Q8 : pas de seuil universel ; notre jeu vise ~8 000 images."))
header(slide, "Veille · Q7–Q8", "Évaluer un GAN : croiser les métriques et le regard")
table(slide, MARGIN, Inches(1.8), Inches(7.6), [
    ["Méthode", "Ce qu'elle mesure", "Limite"],
    ["FID", "Distance entre les statistiques (Inception) des vraies et fausses images", "Biaisé sur petits échantillons"],
    ["KID", "Même idée, estimateur non biaisé", "Dépend aussi d'Inception"],
    ["Inception Score", "Netteté et variété des classes ImageNet", "Inadapté aux peintures"],
    ["Précision / rappel", "Fidélité et couverture, séparément", "Plus coûteux"],
    ["Inspection visuelle", "Artefacts, diversité, mode collapse", "Subjective"],
    ["Plus proches voisins", "Mémorisation des images d'entraînement", "Qualitatif"],
], [1.75, 3.65, 2.2], size=12, row_h=0.6)
box(slide, Inches(8.6), Inches(1.8), Inches(4.1), Inches(4.9))
text(slide, Inches(8.9), Inches(2.05), Inches(3.5), Inches(0.4), "Q8 · Combien de données ?",
     size=20, bold=True, font=HEAD, color=OXBLOOD)
text(slide, Inches(8.9), Inches(2.65), Inches(3.5), Inches(3.9), [
    "Pas de seuil universel : cela dépend de la résolution et de la diversité.",
    "DCGAN : de quelques milliers à des centaines de milliers d'images (CelebA : 200 000).",
    "StyleGAN2-ADA descend à 1 000–2 000 images grâce à l'augmentation adaptative.",
    "Trop peu de données : D sur-apprend, G mémorise.",
    [("Notre choix : ~8 000 images en 64×64.", {"bold": True})],
], size=14, bullets=True, space_after=9)

# Q10-Q11
slide = new_slide(notes="Q10 : avantages et inconvénients. Q11 : variantes avec applications.")
header(slide, "Veille · Q10–Q11", "Des images nettes, au prix d'un entraînement instable")
box(slide, MARGIN, Inches(1.8), Inches(3.0), Inches(4.9))
text(slide, MARGIN + Inches(0.25), Inches(2.0), Inches(2.5), Inches(0.4), "Avantages", size=18,
     bold=True, font=HEAD, color=OXBLOOD)
text(slide, MARGIN + Inches(0.25), Inches(2.55), Inches(2.5), Inches(1.9), [
    "Images nettes et réalistes", "Génération rapide (une passe)", "Pas de vraisemblance à calculer",
    "Non supervisé"], size=14, bullets=True)
text(slide, MARGIN + Inches(0.25), Inches(4.5), Inches(2.5), Inches(0.4), "Inconvénients", size=18,
     bold=True, font=HEAD, color=OXBLOOD)
text(slide, MARGIN + Inches(0.25), Inches(5.0), Inches(2.5), Inches(1.6), [
    "Instable, sensible aux réglages", "Mode collapse", "Évaluation difficile"], size=14, bullets=True)
table(slide, Inches(3.9), Inches(1.8), Inches(8.8), [
    ["Variante", "Idée clé", "Exemple d'application"],
    ["DCGAN", "Réseaux convolutifs", "Images 64×64 (notre baseline)"],
    ["GAN conditionnel", "Ajoute une étiquette à z", "Générer un genre précis"],
    ["Pix2Pix", "Paires d'images", "Croquis → photo, carte → satellite"],
    ["CycleGAN", "Sans paires, cohérence de cycle", "Photo ↔ tableau de Monet"],
    ["StyleGAN", "Style injecté couche par couche", "Visages (thispersondoesnotexist)"],
    ["WGAN-GP", "Distance de Wasserstein", "Entraînement plus stable"],
    ["SRGAN", "Perte perceptuelle", "Super-résolution"],
], [2.1, 3.2, 3.5], size=13, row_h=0.6)

# Q12
slide = new_slide(notes=(
    "Q12 : un CNN est une architecture, un GAN une méthode d'entraînement ; le DCGAN combine les deux."))
header(slide, "Veille · Q12", "CNN, GAN, DCGAN : une architecture et une méthode")
table(slide, MARGIN, Inches(1.8), Inches(12.1), [
    ["", "CNN", "Vanilla GAN (2014)", "DCGAN (2015)"],
    ["Nature", "Architecture de réseau", "Méthode d'entraînement à deux réseaux", "GAN dont G et D sont des CNN"],
    ["Tâche typique", "Classer (discriminant)", "Générer", "Générer des images"],
    ["Couches", "Convolutions, pooling", "Perceptrons multicouches (denses)", "Convolutions à pas, convolutions transposées dans G"],
    ["Normalisation", "Variable", "Aucune", "BatchNorm dans G et D"],
    ["Activations", "ReLU", "ReLU / sigmoïde", "ReLU (G), LeakyReLU (D), tanh en sortie"],
    ["Limite", "Ne génère pas seul", "Instable sur images complexes", "Résolution limitée (~64×64)"],
], [2.0, 2.8, 3.3, 4.0], size=14, row_h=0.62)
footer(slide, "Radford, Metz, Chintala, Unsupervised Representation Learning with DCGANs, 2015")

# Q13-Q15
slide = new_slide(notes=(
    "Q13 : ChatGPT repose sur GPT, un transformer autorégressif d'OpenAI. Q14 : dangers. "
    "Q15 : notre position sur l'intérêt de générer des données."))
header(slide, "Veille · Q13–Q15", "D'autres modèles génératifs, et leurs risques")
cards = [
    ("Q13 · ChatGPT", [
        [("Modèle : ", {"bold": True}), ("GPT, Generative Pre-trained Transformer.", {})],
        [("Famille : ", {"bold": True}), ("transformers, grands modèles de langage autorégressifs (prédiction du mot suivant), affinés par retour humain (RLHF).", {})],
        [("Développé par ", {"bold": True}), ("OpenAI.", {})],
    ]),
    ("Q14 · Dangers", [
        "Deepfakes : usurpation d'identité, fraude, chantage",
        "Désinformation à grande échelle",
        "Biais des données reproduits",
        "Droits des artistes, imitation de style",
        "Mémorisation de données privées",
    ]),
    ("Q15 · Pourquoi générer ?", [
        "Compléter des données rares (imagerie médicale)",
        "Données synthétiques anonymes",
        "Outil de création pour les artistes",
        [("Utile si le contenu synthétique reste identifié comme tel.", {"bold": True})],
    ]),
]
for k, (title, body) in enumerate(cards):
    x = MARGIN + k * (cw + Inches(0.3))
    box(slide, x, Inches(1.8), cw, Inches(4.9))
    text(slide, x + Inches(0.3), Inches(2.05), cw - Inches(0.6), Inches(0.4), title, size=20,
         bold=True, font=HEAD, color=OXBLOOD)
    text(slide, x + Inches(0.3), Inches(2.65), cw - Inches(0.6), Inches(3.9), body, size=16,
         bullets=True, space_after=12)

# ================================================================ Catalogue
divider("02", "Le catalogue", "Analyse exploratoire des métadonnées, avant tout téléchargement",
        "Deuxième partie : l'analyse du catalogue, faite sans télécharger d'images.")

slide = new_slide(notes=(
    "Les chiffres clés du catalogue. Les inconnues sont codées par un tiret, pas par des cases vides."))
header(slide, "Le catalogue", "Un catalogue riche, à 61 % de peintures")
stats = [
    (fmt(EDA["n_catalogue"]), "œuvres au catalogue"),
    (fmt(n_paint), f"peintures ({pct(n_paint / EDA['n_catalogue'])})"),
    (fmt(EDA["n_artists"]), "artistes peintres"),
    (pct(EDA["share_known_year"]), "des peintures datées"),
]
sw = Inches(2.85)
for k, (value, label) in enumerate(stats):
    x = MARGIN + k * (sw + Inches(0.23))
    box(slide, x, Inches(1.9), sw, Inches(2.0))
    text(slide, x, Inches(2.1), sw, Inches(0.9), value, size=44, bold=True, font=HEAD,
         color=OXBLOOD, align=PP_ALIGN.CENTER)
    text(slide, x, Inches(3.1), sw, Inches(0.5), label, size=15, color=GREY, align=PP_ALIGN.CENTER)
text(slide, MARGIN, Inches(4.4), Inches(12), Inches(2.3), [
    [("URL unique pour chaque œuvre : ", {"bold": True}), ("elle sert d'identifiant pour relier chaque image à sa source.", {})],
    [("Aucune case vide, ", {"bold": True}), ("mais les inconnues sont codées par « - » (12 % des dates de peintures).", {})],
    [("Champs en texte libre : ", {"bold": True}),
     ("dates (« c. 1530 », « 1648-50 »), dimensions (« Oil on canvas, 68 x 95 cm ») ; il a fallu les analyser pour en tirer des nombres.", {})],
], size=15, bullets=True, space_after=10)
footer(slide, "Source : Web Gallery of Art (wga.hu) · notebook 01_eda")

slide = new_slide(notes=(
    "Le déséquilibre des genres est le premier risque : un GAN entraîné au hasard ferait surtout des scènes religieuses."))
header(slide, "Le catalogue", f"Le religieux domine : {pct(types['religious'] / n_paint)} des peintures")
order = types.sort_values()
data = CategoryChartData()
data.categories = [GENRE_FR[g] for g in order.index]
data.add_series("Part des peintures", [v / n_paint for v in order.values])
gf = slide.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, MARGIN, Inches(1.7), Inches(7.6), Inches(5.2), data)
chart = gf.chart
style_chart(chart)
chart.has_title = False
chart.has_legend = False
plot = chart.plots[0]
plot.gap_width = 45
plot.has_data_labels = True
plot.data_labels.number_format = '0%'
plot.data_labels.number_format_is_linked = False
plot.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
plot.data_labels.font.size = Pt(12)
series = plot.series[0]
series.format.fill.solid()
series.format.fill.fore_color.rgb = STONE
point = series.points[len(order) - 1]
point.format.fill.solid()
point.format.fill.fore_color.rgb = OXBLOOD
chart.value_axis.visible = False
chart.value_axis.has_major_gridlines = False
chart.category_axis.tick_labels.font.size = Pt(13)
chart.category_axis.format.line.fill.background()
text(slide, Inches(8.7), Inches(2.0), Inches(4.0), Inches(4.5), [
    [("Conséquence : ", {"bold": True, "color": OXBLOOD}),
     ("un échantillon tiré au hasard produirait surtout des madones et des scènes bibliques.", {})],
    [("Les natures mortes, intérieurs et scènes de genre sont minoritaires.", {})],
    [("Décision : ", {"bold": True, "color": OXBLOOD}), ("quotas équilibrés par genre lors de la sélection.", {})],
], size=16, space_after=14)

slide = new_slide(notes=(
    "Le genre dépend fortement de l'époque : rééquilibrer les genres rééquilibre aussi les époques."))
header(slide, "Le catalogue", "Le genre suit l'époque : du religieux au paysage")
slide.shapes.add_picture(str(FIG / "07_type_by_century.png"), MARGIN, Inches(1.65), height=Inches(5.3))
text(slide, Inches(9.4), Inches(2.0), Inches(3.3), Inches(4.5), [
    [("99 % ", {"bold": True, "color": OXBLOOD, "size": 28, "font": HEAD})],
    [("de sujets religieux au XIIIe siècle, 3 % au XIXe.", {})],
    [("34 % ", {"bold": True, "color": OXBLOOD, "size": 28, "font": HEAD})],
    [("de paysages au XIXe siècle.", {})],
    [("Couvrir tous les genres, c'est aussi couvrir toutes les époques.", {"bold": True})],
], size=15, space_after=6)
footer(slide, "Siècle calculé à partir de la date de l'œuvre (champ DATE)")

slide = new_slide(notes=(
    "Quatre pièges repérés pendant l'analyse, chacun avec sa conséquence sur la sélection."))
header(slide, "Le catalogue", "Quatre pièges repérés avant de télécharger")
traps = [
    (pct(EDA["n_detail_views"] / n_paint), "de vues de détail",
     "Des recadrages d'une autre œuvre : quasi-doublons et risque de fuite entre entraînement et évaluation.",
     "Exclues, découpage par œuvre"),
    (pct(EDA["techniques"]["fresco"] / n_paint), "de fresques",
     "Photos de murs et de voûtes, éclairage inégal, formats hétérogènes.",
     "Exclues"),
    (pct(1 - EDA["share_year_inside_timeframe"]), "de périodes incohérentes",
     "La colonne TIMEFRAME ne correspond pas à la date de l'œuvre.",
     "Date relue dans DATE"),
    ("0,85", "de ratio médian",
     "Formats surtout verticaux ; un GAN génère des carrés. Recadrer au centre perd 24 % de la surface en médiane.",
     "Recadrage carré, ratios extrêmes exclus"),
]
tw = Inches(2.85)
for k, (value, label, body, decision) in enumerate(traps):
    x = MARGIN + k * (tw + Inches(0.23))
    box(slide, x, Inches(1.8), tw, Inches(4.9))
    text(slide, x + Inches(0.25), Inches(2.0), tw - Inches(0.5), Inches(0.8), value, size=38,
         bold=True, font=HEAD, color=OXBLOOD)
    text(slide, x + Inches(0.25), Inches(2.85), tw - Inches(0.5), Inches(0.4), label, size=15, bold=True)
    text(slide, x + Inches(0.25), Inches(3.4), tw - Inches(0.5), Inches(2.0), body, size=14, color=GREY)
    text(slide, x + Inches(0.25), Inches(5.65), tw - Inches(0.5), Inches(0.8),
         [[("→ ", {"bold": True, "color": OXBLOOD}), (decision, {"bold": True})]], size=14)

# ================================================================ Jeu d'images
divider("03", "Le jeu d'images", "Sélection, scraping et prétraitement",
        "Troisième partie : comment on passe du catalogue à des images prêtes pour l'entraînement.")

slide = new_slide(notes=(
    "Entonnoir de sélection, puis tirage de 8 000 images avec quotas équilibrés par genre et 30 images "
    "maximum par artiste. Découpage 90/10 par œuvre."))
header(slide, "Le jeu d'images", f"{fmt(n_selected)} peintures choisies, genres rééquilibrés")
funnel = EDA["selection_funnel"]
steps = [
    (funnel["peintures"], "peintures"),
    (funnel["sans vues de détail"], "sans vues de détail"),
    (funnel["+ sans fresques"], "sans fresques"),
    (funnel["+ ratio connu dans [0.5, 2] ou inconnu"], "formats raisonnables"),
    (n_selected, "sélection finale"),
]
for k, (value, label) in enumerate(steps):
    y = Inches(1.8) + k * Inches(0.98)
    width = Inches(4.6) * (0.45 + 0.55 * value / steps[0][0])
    last = k == len(steps) - 1
    box(slide, MARGIN, y, width, Inches(0.8), fill=OXBLOOD if last else CARD)
    text(slide, MARGIN + Inches(0.2), y, width - Inches(0.3), Inches(0.8),
         [[(fmt(value), {"bold": True, "size": 20, "font": HEAD}), ("  " + label, {})]], size=14,
         color=WHITE if last else INK, anchor=MSO_ANCHOR.MIDDLE)
sel_types = MANIFEST["TYPE"].value_counts(normalize=True)
order = pd.Series(EDA["types"]).sort_values(ascending=True).index
data = CategoryChartData()
data.categories = [GENRE_FR[g] for g in order]
data.add_series("Catalogue", [EDA["types"][g] / n_paint for g in order])
data.add_series("Sélection", [float(sel_types.get(g, 0)) for g in order])
gf = slide.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(5.7), Inches(1.65), Inches(7.0), Inches(5.3), data)
chart = gf.chart
style_chart(chart)
chart.has_title = False
chart.has_legend = True
chart.legend.position = XL_LEGEND_POSITION.TOP
chart.legend.include_in_layout = False
chart.legend.font.size = Pt(13)
plot = chart.plots[0]
plot.gap_width = 40
plot.overlap = -10
plot.has_data_labels = True
plot.data_labels.number_format = '0%'
plot.data_labels.number_format_is_linked = False
plot.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
plot.data_labels.font.size = Pt(11)
for s, color in zip(plot.series, (STONE, OXBLOOD)):
    s.format.fill.solid()
    s.format.fill.fore_color.rgb = color
chart.value_axis.visible = False
chart.value_axis.has_major_gridlines = False
chart.category_axis.tick_labels.font.size = Pt(13)
chart.category_axis.format.line.fill.background()
footer(slide, "Une œuvre par ligne, 30 peintures max. par artiste, découpage 90/10 par œuvre (pas de fuite)")

slide = new_slide(notes=(
    "Le pipeline de scraping est reprenable. Difficulté rencontrée : le premier téléchargement s'est bloqué "
    "car le serveur envoyait certaines images au compte-gouttes ; corrigé par un délai total par image."))
header(slide, "Le jeu d'images", "Un scraping reprenable, image par image")
flow = [
    ("Manifeste", "1 ligne par image, statut pending"),
    ("Téléchargement", "3 essais, 45 s max. par image"),
    ("Validation", "Image lisible, ≥ 256 px, ratio réel"),
    ("Recadrage", "Carré central, 256×256, RGB"),
    ("Statut", "processed, failed ou skipped"),
]
fw = Inches(2.1)
for k, (title, body) in enumerate(flow):
    x = MARGIN + k * (fw + Inches(0.4))
    box(slide, x, Inches(1.85), fw, Inches(1.6), fill=INK if k in (0, 4) else CARD)
    dark = k in (0, 4)
    text(slide, x + Inches(0.15), Inches(2.0), fw - Inches(0.3), Inches(0.4), title, size=16, bold=True,
         color=GOLD if dark else OXBLOOD, font=HEAD, align=PP_ALIGN.CENTER)
    text(slide, x + Inches(0.15), Inches(2.5), fw - Inches(0.3), Inches(0.9), body, size=13,
         color=WHITE if dark else INK, align=PP_ALIGN.CENTER)
    if k < len(flow) - 1:
        arrow(slide, x + fw + Inches(0.07), Inches(2.5), Inches(0.26), Inches(0.3))
progress = [
    (fmt(n_processed), "images traitées"),
    (fmt(n_failed), "en échec (retentées)"),
    (fmt(n_skipped), "écartées (doublons, formats)"),
    (fmt(n_pending), "en attente"),
]
pw = Inches(2.85)
for k, (value, label) in enumerate(progress):
    x = MARGIN + k * (pw + Inches(0.23))
    text(slide, x, Inches(3.85), pw, Inches(0.7), value, size=36, bold=True, font=HEAD,
         color=OXBLOOD if k == 0 else INK)
    text(slide, x, Inches(4.55), pw, Inches(0.4), label, size=14, color=GREY)
text(slide, MARGIN, Inches(5.25), Inches(12.1), Inches(1.4), [
    [("Difficulté : ", {"bold": True, "color": OXBLOOD}),
     ("le premier téléchargement s'est bloqué sans erreur. Certaines images arrivaient au compte-gouttes "
      "et la limite de temps ne portait que sur chaque lecture réseau. Correction : un délai total par image, "
      "couvert par un test.", {})],
], size=14)
footer(slide, ("Téléchargement terminé" if download_done else f"Téléchargement en cours, état au {stamp}")
       + " · stockage : ~34 Ko par image")

slide = new_slide(notes=(
    "Question 5 de l'énoncé : ce que montrent les images, et les conséquences sur le prétraitement."))
header(slide, "Le jeu d'images · Q5", "Ce que montrent les images")
genres = ["portrait", "landscape", "still-life", "religious", "genre"]
side, gap = Inches(0.92), Inches(0.08)
for r, genre in enumerate(genres):
    y = Inches(1.7) + r * (side + gap)
    text(slide, MARGIN, y, Inches(1.5), side, GENRE_FR[genre], size=13, bold=True, color=GREY,
         anchor=MSO_ANCHOR.MIDDLE)
    for c, path in enumerate(processed_images(5, seed=r, genre=genre)):
        slide.shapes.add_picture(str(path), MARGIN + Inches(1.5) + c * (side + gap), y, side, side)
text(slide, Inches(7.4), Inches(1.8), Inches(5.3), Inches(5.0), [
    [("Tons sombres et bruns : ", {"bold": True}), ("vernis anciens, fonds noirs des portraits et natures mortes.", {})],
    [("Cadres visibles ", {"bold": True}), ("sur certaines photos, ou mur autour de l'œuvre.", {})],
    [("Sujet souvent centré : ", {"bold": True}), ("le recadrage carré central le conserve.", {})],
    [("Reproductions noir et blanc ", {"bold": True}), ("à détecter et exclure.", {})],
    [("Résolution d'origine ", {"bold": True}), ("largement supérieure à 256 px.", {})],
], size=15, bullets=True, space_after=12)
footer(slide, "Images réelles téléchargées, après recadrage carré (256×256)")

slide = new_slide(notes="Les décisions de prétraitement, chacune justifiée par une observation.")
header(slide, "Le jeu d'images", "Un prétraitement justifié par les observations")
table(slide, MARGIN, Inches(1.8), Inches(12.1), [
    ["Observation", "Décision", "Pourquoi"],
    ["Formats verticaux et horizontaux", "Recadrage carré central", "Le DCGAN génère des images carrées"],
    ["Ressources limitées, pas de GPU local", "Entraînement en 64×64", "Baseline rapide ; 128×128 possible sans retélécharger"],
    ["Tonalités sombres", "Pixels normalisés dans [-1, 1]", "Sortie tanh du générateur ; on garde les couleurs d'origine"],
    ["Reproductions noir et blanc", "Exclues (saturation faible)", "Sinon le GAN apprend à produire du gris"],
    ["Quasi-doublons", "Une image par paire (empreinte aHash)", "Évite mémorisation et fuite vers l'évaluation"],
    ["Compositions sans sens gauche / droite", "Retournement horizontal aléatoire", "Double la variété ; pas de retournement vertical"],
], [3.6, 3.9, 4.6], size=14, row_h=0.66)
footer(slide, "Seuils et effectifs exacts : notebook 02_dataset, à valider sur le jeu complet")

# ================================================================ Modele
divider("04", "Le modèle", "Baseline DCGAN 64×64",
        "Quatrième partie : le modèle. L'entraînement n'a pas encore été lancé.")

slide = new_slide(notes=(
    "Architecture prévue, d'après le papier DCGAN. Rien n'est encore entraîné : ces valeurs sont un point "
    "de départ à valider."))
header(slide, "Le modèle · prévu", "Architecture DCGAN retenue comme baseline")
gen = [("z", "100", 0.35), ("4×4", "512", 0.55), ("8×8", "256", 0.85), ("16×16", "128", 1.2),
       ("32×32", "64", 1.6), ("64×64", "3", 2.0)]
x = MARGIN
text(slide, MARGIN, Inches(1.75), Inches(7), Inches(0.4), "Générateur : convolutions transposées",
     size=16, bold=True, font=HEAD, color=OXBLOOD)
for k, (res, ch, size) in enumerate(gen):
    h = Inches(size)
    y = Inches(3.3) - h / 2
    box(slide, x, y, Inches(0.75), h, fill=OXBLOOD if k == len(gen) - 1 else STONE, radius=0.08)
    text(slide, x - Inches(0.15), Inches(4.45), Inches(1.05), Inches(0.5),
         [[(res, {"bold": True})], [(ch + (" (RGB)" if k == len(gen) - 1 else " canaux"), {"size": 11, "color": GREY})]],
         size=13, align=PP_ALIGN.CENTER, space_after=0)
    if k < len(gen) - 1:
        arrow(slide, x + Inches(0.82), Inches(3.15), Inches(0.28), Inches(0.28))
    x += Inches(1.2)
text(slide, MARGIN, Inches(5.25), Inches(7.1), Inches(1.4), [
    [("Générateur : ", {"bold": True}), ("BatchNorm + ReLU, sortie tanh.", {})],
    [("Discriminateur : ", {"bold": True}), ("architecture miroir (convolutions à pas), LeakyReLU 0,2, sortie sigmoïde.", {})],
], size=14, bullets=True)
box(slide, Inches(8.1), Inches(1.8), Inches(4.6), Inches(4.9))
text(slide, Inches(8.4), Inches(2.05), Inches(4.0), Inches(0.4), "Hyperparamètres de départ", size=18,
     bold=True, font=HEAD, color=OXBLOOD)
text(slide, Inches(8.4), Inches(2.65), Inches(4.0), Inches(3.9), [
    [("Perte : ", {"bold": True}), ("entropie croisée binaire", {})],
    [("Optimiseur : ", {"bold": True}), ("Adam, lr 0,0002, β1 = 0,5", {})],
    [("Lot : ", {"bold": True}), ("128 images", {})],
    [("Dimension de z : ", {"bold": True}), ("100", {})],
    [("Matériel : ", {"bold": True}), ("GPU Colab (pas de GPU NVIDIA local)", {})],
    [("Valeurs du papier DCGAN, à valider sur nos données.", {"italic": True, "color": GREY})],
], size=14, bullets=True, space_after=9)
footer(slide, "Statut : prévu, non entraîné")

slide = new_slide(notes=(
    "Les résultats seront ajoutés après l'entraînement. Le protocole d'évaluation est déjà fixé."))
header(slide, "Le modèle · résultats", "Résultats : entraînement à venir")
box(slide, MARGIN, Inches(1.8), Inches(5.6), Inches(4.9), fill=WHITE, line=STONE)
text(slide, MARGIN, Inches(1.8), Inches(5.6), Inches(4.9),
     [[("Emplacement des grilles d'images générées", {"bold": True})],
      [("à compléter après l'entraînement", {"color": GREY})]],
     size=16, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
text(slide, Inches(6.8), Inches(1.9), Inches(5.9), Inches(0.4), "Protocole d'évaluation prévu", size=18,
     bold=True, font=HEAD, color=OXBLOOD)
text(slide, Inches(6.8), Inches(2.5), Inches(5.9), Inches(4.2), [
    "Grilles générées à partir des mêmes vecteurs z à chaque époque, pour suivre la progression",
    "Courbes de pertes de G et D, croisées avec les images (les pertes seules ne suffisent pas)",
    "FID et KID entre images générées et jeu d'évaluation (~10 %, jamais vu à l'entraînement)",
    "Plus proches voisins dans le jeu d'entraînement, pour détecter la mémorisation",
    "Checkpoints complets (poids, optimiseurs, époque, graines) pour reprendre sur Colab",
], size=15, bullets=True, space_after=10)
footer(slide, "Statut : non commencé")

# ================================================================ Equipe
slide = new_slide(notes="À compléter par l'équipe : outils, rituels, répartition des tâches.")
header(slide, "Équipe", "Organisation de l'équipe")
team = [
    ("Outils", ["Dépôt GitHub : [à compléter]", "Suivi des tâches (Trello / Azure DevOps) : [à compléter]"]),
    ("Rituels", ["Points d'équipe : [fréquence à compléter]", "Revues de code : [à compléter]"]),
    ("Répartition", ["[Membre 1] : [rôle]", "[Membre 2] : [rôle]", "[Membre 3] : [rôle]"]),
]
for k, (title, body) in enumerate(team):
    x = MARGIN + k * (cw + Inches(0.3))
    box(slide, x, Inches(1.8), cw, Inches(4.9))
    circle(slide, x + Inches(0.3), Inches(2.1), Inches(0.6), str(k + 1))
    text(slide, x + Inches(1.1), Inches(2.18), cw - Inches(1.4), Inches(0.45), title, size=20,
         bold=True, font=HEAD, color=OXBLOOD)
    text(slide, x + Inches(0.3), Inches(3.0), cw - Inches(0.6), Inches(3.5), body, size=15,
         bullets=True, space_after=10)

# ================================================================ Conclusion
slide = new_slide(dark=True, notes=(
    "Bilan : analyse et pipeline de données terminés et testés ; le GAN reste à entraîner."))
text(slide, MARGIN, Inches(0.45), Inches(12), Inches(0.3), "CONCLUSION", size=12, bold=True, color=GOLD)
text(slide, MARGIN, Inches(0.75), Inches(12), Inches(0.8), "Des données solides avant le premier GAN",
     size=32, bold=True, font=HEAD, color=WHITE)
cols = [
    ("Acquis", [
        "Catalogue analysé, pièges identifiés",
        f"{fmt(n_selected)} peintures sélectionnées, genres équilibrés",
        "Scraping reprenable et testé (25 tests)",
    ]),
    ("Difficultés", [
        "Serveur lent et irrégulier : blocage puis correctif",
        "Données hétérogènes : cadres, noir et blanc",
        "Fort déséquilibre des genres",
    ]),
    ("Prochaines étapes", [
        "Terminer le téléchargement et le contrôle des images",
        "Entraîner la baseline DCGAN sur Colab",
        "Évaluer (FID, KID, visuel), puis l'application Streamlit",
    ]),
]
for k, (title, body) in enumerate(cols):
    x = MARGIN + k * (cw + Inches(0.3))
    text(slide, x, Inches(2.0), cw, Inches(0.5), title, size=22, bold=True, font=HEAD, color=GOLD)
    text(slide, x, Inches(2.7), cw, Inches(3.5), body, size=16, color=WHITE, bullets=True, space_after=12)
strip = processed_images(9, seed=21)
side = Inches(1.25)
for k, path in enumerate(strip):
    slide.shapes.add_picture(str(path), MARGIN + k * (side + Inches(0.11)), Inches(5.55), side, side)

prs.core_properties.title = "GANisme"
prs.core_properties.subject = "Générer des peintures avec un GAN"
OUT.parent.mkdir(parents=True, exist_ok=True)
prs.save(OUT)
print(f"{OUT} : {len(prs.slides)} diapositives ; images traitees : {n_processed} ; etat {stamp}")
